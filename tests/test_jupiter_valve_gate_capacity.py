"""Jupiter's template holds 12 active gates in rows 66..77, not nine.

Fixtures are synthetic recipes and never read or alter production Moldes.xlsx.
They exercise the real parser, family-scoped resolver, planner and template.
"""
from __future__ import annotations

from copy import copy
from decimal import Decimal
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from app.catalogs import DataCatalog, MappingCatalog
from app.errors import ProcessWarning, SourceDataError
from app.input_policy import JUPITER_GROUP
from app.jupiter_xml_parser import JupiterXmlParser
from app.main import _hrs_total_zones_text
from app.plan import PlanBuilder
from app.resolver import DataResolver

ROOT = Path(__file__).resolve().parents[1]
LAYOUT = 'JUPITER_TWOSHOT_1080_V1'
SOURCES = {'PARAM', 'MOLD', 'MANUAL'}
GATES = (1, 3, 9, 13, 14, 17, 19, 21, 22, 23, 37, 48)
COLUMNS = {
    'GATE_NO': 'AV', 'START_USE': 'AX', 'START_CON': 'BA',
    'ON_STROKE': 'BE', 'START_DELAY': 'BI', 'STOP_CON': 'BN',
    'OFF_STROKE': 'BR', 'STOP_DELAY': 'BV', 'HOLDING_STATE': 'CA',
    'HOLDING_DELAY': 'CD', 'HOLDING_TIME': 'CH', 'INJECTION_NO': 'CM',
}


@pytest.fixture(scope='module')
def catalogs():
    return DataCatalog(ROOT / 'data/Data.xlsx'), MappingCatalog(ROOT / 'data/Mapeo.xlsx')


def _recipe(path, machine, gates, data, *, hold=True):
    root = ET.Element('HMI_Data', {'Version': '1.2'})
    group = ET.SubElement(root, 'VarGroup', {'Name': 'VG_MoldData'})
    # Valid profile counts unrelated to the sequencer under test.
    values = {
        'Mold1.sv_MoldFwdProfVisSrc.Profile.iNoOfPoints': 5,
        'Mold1.sv_MoldBwdProfVis.Profile.iNoOfPoints': 1,
        'Injection1.sv_InjectProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_InjectProfVis.Profile.iNoOfPoints': 1,
        'Injection1.sv_HoldProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_HoldProfVis.Profile.iNoOfPoints': 1,
        'RotaryTable1.sv_nRotaryTableMode': 2,
    }
    for core in (1, 2):
        values.update({
            f'Core{core}.sv_CoreMode.CoreType': 1,
            f'Core{core}.sv_CoreMode.CoreControlIn': 2,
            f'Core{core}.sv_CoreMode.CoreControlOut': 1,
            f'CentralCoordination1.sv_CoreData[{core}].InMode': 0,
            f'CentralCoordination1.sv_CoreData[{core}].OutMode': 2,
        })
    constants = data.machine_constants[machine]
    for i, gate in enumerate(gates, 1):
        unit = 1 if i % 2 else 2
        area = Decimal(str(constants[f'U{unit}_SCREW_VOLUME_CM3'])) / Decimal(str(constants[f'U{unit}_SCREW_STROKE_MM']))
        key = f'ValveGate1.sv_ValveGateData.ValveGateDataArray[{gate}]'
        suffix = '_Inj2' if unit == 2 else ''
        fields = {
            'bUsed': True, 'injection': unit - 1,
            'OpenMode': i % 2, 'CloseMode': (i + 1) % 2,
            f'rStartPos{suffix}': area * (20 + i),
            f'rStopPos{suffix}': area * (1 + i),
            'dStartDelay': i * 100_000, 'dStopDelay': i * 200_000,
            'bActivOnHold': hold,
            # Nonzero residuals must be hidden while HOLD is off.
            'dHoldDelay': i * 300_000, 'dHoldActTime': i * 400_000,
        }
        values.update({f'{key}.{field}': value for field, value in fields.items()})
    for name, value in values.items():
        element = ET.SubElement(group, 'Variable', {'Name': name})
        ET.SubElement(element, 'Value').text = str(value).lower() if isinstance(value, bool) else str(value)
    ET.ElementTree(root).write(path, encoding='utf-8', xml_declaration=True)
    return path


def _resolve(catalogs, machine, path, *, strict=True):
    data, mapping = catalogs
    parser = JupiterXmlParser(machine=machine, config=data.config,
                              machine_constants=data.machine_constants,
                              strict_template_limits=strict)
    records = parser.parse(path)
    resolver = DataResolver(data)
    required = mapping.required_keys_for(JUPITER_GROUP, LAYOUT, enabled_sources=SOURCES)
    values = resolver.resolve_param(records, required_keys=required - {'HRSTotalZonesText'},
                                    template_group=JUPITER_GROUP)
    assert resolver.parameter_warnings == []
    values['HRSTotalZonesText'] = _hrs_total_zones_text(values)
    return parser, values


def _build(catalogs, machine, values):
    data, mapping = catalogs
    return PlanBuilder(mapping, data).build(JUPITER_GROUP, LAYOUT, values,
        {'MachineNumber': machine, 'MoldNumber': 'I-9000'}, enabled_sources=SOURCES)


def _apply(plan):
    book = load_workbook(ROOT / 'plantillas/Haitian Jupiter TwoShot_1080.xlsx')
    ws = book.active
    # G20 is also a process field: snapshot source fills before writes, as the writer does.
    fills = {op.ref_cell: copy(ws[op.ref_cell].fill) for op in plan.operations if op.kind == 'COPY_FILL'}
    for op in plan.operations:
        if op.kind == 'COPY_FILL':
            ws[op.cell].fill = copy(fills[op.ref_cell])
        elif op.kind == 'WRITE':
            ws[op.cell] = op.value
        elif op.kind == 'CLEAR':
            ws[op.cell] = None
        else:
            raise AssertionError(op.kind)
    return book, ws, fills


def test_parser_data_and_mapping_agree_on_twelve_slots(catalogs):
    data, mapping = catalogs
    settings = mapping.valve_gate_settings(JUPITER_GROUP, LAYOUT)
    assert JupiterXmlParser.MAX_VALVE_GATE_ACTIVE == 12
    assert int(data.cfg('JUPITER_VALVE_GATE_MAX_ACTIVE')) == 12
    assert int(settings['TEMPLATE_MAX_ACTIVE_GATES']) == 12
    assert int(settings['CONTROLLER_MAX_GATES']) == 48
    assert settings['NUMBERING_POLICY'] == 'DIRECT'
    assert settings['OVER_LIMIT_ACTION'] in {'WARNING', 'ERROR', 'BLOCK'}


def test_all_12_fields_use_correct_merged_anchors_in_rows_66_through_77(catalogs):
    _, mapping = catalogs
    rows = mapping.valve_gate_layout(JUPITER_GROUP, LAYOUT)
    expected = {(f'SLOT{slot}', field): f'{col}{65 + slot}'
                for slot in range(1, 13) for field, col in COLUMNS.items()}
    actual = {(r['SLOT'], r['FIELD']): r['CELL'] for r in rows}
    assert len(rows) == len(actual) == 144
    assert actual == expected
    # Both visible autofilters and their hidden Excel names must include new slots.
    with ZipFile(ROOT / 'data/Mapeo.xlsx') as archive:
        workbook_xml = archive.read('xl/workbook.xml').decode('utf-8')
        assert "'ValveGateLayout'!$A$1:$G$525" in workbook_xml
        assert "'ValveGateLogic'!$A$1:$E$78" in workbook_xml
    book = load_workbook(ROOT / 'plantillas/Haitian Jupiter TwoShot_1080.xlsx')
    try:
        ws = book.active
        anchors = {ws.cell(r.min_row, r.min_col).coordinate for r in ws.merged_cells.ranges}
        assert set(actual.values()) <= anchors
        assert not ({'A75', 'A76'} & set(actual.values()))
    finally:
        book.close()


@pytest.mark.parametrize('machine', ['308B', '309'])
@pytest.mark.parametrize('count', [0, 1, 4, 9, 10, 11, 12])
def test_full_pipeline_accepts_capacity_and_shades_only_unused_rows(catalogs, tmp_path, machine, count):
    gates = GATES[:count]
    path = _recipe(tmp_path / 'recipe.xml', machine, gates, catalogs[0])
    parser, values = _resolve(catalogs, machine, path)
    assert parser.process_warnings == []
    assert parser.active_valve_gates == gates
    plan = _build(catalogs, machine, values)
    assert plan.active_valve_gates == gates
    book, ws, fills = _apply(plan)
    try:
        for slot in range(1, 13):
            row = 65 + slot
            if slot <= count:
                assert ws[f'AV{row}'].value == gates[slot - 1]
                assert ws[f'AX{row}'].value == '\u2713'
                assert ws[f'BA{row}'].value == ('TIME' if slot % 2 else 'POSITION')
                assert ws[f'BN{row}'].value == ('POSITION' if slot % 2 else 'TIME')
                assert float(ws[f'BE{row}'].value) == pytest.approx(20 + slot)
                assert float(ws[f'BR{row}'].value) == pytest.approx(1 + slot)
                for col, factor in [('BI', .1), ('BV', .2), ('CD', .3), ('CH', .4)]:
                    assert float(ws[f'{col}{row}'].value) == pytest.approx(slot * factor)
                assert ws[f'CA{row}'].value == '\u2713'
                assert ws[f'CM{row}'].value == ('Inyection 1' if slot % 2 else 'Inyection 2')
                for col in COLUMNS.values():
                    assert copy(ws[f'{col}{row}'].fill) == fills['G17']
            else:
                for col in COLUMNS.values():
                    assert ws[f'{col}{row}'].value is None
                    assert copy(ws[f'{col}{row}'].fill) == fills['G20']
        # Preserve template-owned content beside the three additional slots.
        assert ws['A75'].value == 'OBSERVACIONES'
        assert 'NO SON AUDITABLES' in ws['A76'].value
        for col in ('B', 'Z', 'BA', 'BY'):
            assert [ws[f'{col}{r}'].value for r in range(26, 30)] == [1, 2, 3, 'End']
        assert len(ws.merged_cells.ranges) == 730
        assert not any(op.kind == 'WRITE' and op.cell in {f'AV{r}' for r in range(78, 90)} for op in plan.operations)
    finally:
        book.close()


@pytest.mark.parametrize('machine', ['308B', '309'])
def test_thirteenth_active_gate_is_rejected_by_parser(catalogs, tmp_path, machine):
    gates = tuple(sorted((*GATES, 47)))
    path = _recipe(tmp_path / 'recipe13.xml', machine, gates, catalogs[0])
    with pytest.raises(SourceDataError, match='maximo 12'):
        _resolve(catalogs, machine, path)


@pytest.mark.parametrize('machine', ['308B', '309'])
def test_planner_independently_blocks_thirteen_even_with_non_strict_parser(catalogs, tmp_path, machine):
    path = _recipe(tmp_path / 'recipe13.xml', machine, tuple(sorted((*GATES, 47))), catalogs[0])
    parser, values = _resolve(catalogs, machine, path, strict=False)
    assert any('maximo 12' in message for message in parser.process_warnings)
    with pytest.raises(ProcessWarning, match='solo admite 12'):
        _build(catalogs, machine, values)


@pytest.mark.parametrize('machine', ['308B', '309'])
def test_new_slots_hide_residual_hold_values_when_hold_disabled(catalogs, tmp_path, machine):
    path = _recipe(tmp_path / 'hold_off.xml', machine, GATES, catalogs[0], hold=False)
    _, values = _resolve(catalogs, machine, path)
    book, ws, fills = _apply(_build(catalogs, machine, values))
    try:
        for row in (75, 76, 77):
            assert ws[f'AV{row}'].value == GATES[row - 66]
            assert ws[f'CA{row}'].value is None
            for col in ('CD', 'CH'):
                assert ws[f'{col}{row}'].value is None
                assert copy(ws[f'{col}{row}'].fill) == fills['G20']
            assert copy(ws[f'AV{row}'].fill) == fills['G17']
    finally:
        book.close()
