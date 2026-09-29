from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.catalogs import MappingCatalog
from app.errors import ProcessWarning
from app.plan import PlanBuilder


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _activate_310b_mapping(tmp_path: Path) -> Path:
    source = PROJECT_ROOT / "data" / "Mapeo.xlsx"
    target = tmp_path / "Mapeo_310B_active.xlsx"
    wb = load_workbook(source)

    ws = wb["Machines"]
    headers = {cell.value: cell.column for cell in ws[1]}
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, headers["MACHINE"]).value).strip().upper() == "310B":
            ws.cell(row, headers["ACTIVE"]).value = True
            break
    else:
        raise AssertionError("310B is missing from Machines")

    ws = wb["Templates"]
    headers = {cell.value: cell.column for cell in ws[1]}
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, headers["LAYOUT ID"]).value).strip() == "ZERES_GENIII_1080_V1":
            ws.cell(row, headers["ACTIVE"]).value = True
            break
    else:
        raise AssertionError("ZERES_GENIII_1080_V1 is missing from Templates")

    wb.save(target)
    return target


def test_machine_specific_layout_allows_500_and_1080_in_same_family(tmp_path):
    mapping = MappingCatalog(_activate_310b_mapping(tmp_path))

    group_500, template_500 = mapping.resolve_machine("105A")
    group_1080, template_1080 = mapping.resolve_machine("310B")

    assert group_500 == "HAITIAN_ZE_III"
    assert template_500.layout_id == "ZERES_GENIII_V1"
    assert template_500.template_file == "Haitian Zeres Gen III_500.xlsx"

    assert group_1080 == "HAITIAN_ZE_III"
    assert template_1080.layout_id == "ZERES_GENIII_1080_V1"
    assert template_1080.template_file == "Haitian Zeres Gen III_1080.xlsx"


def _core_values(mapping: MappingCatalog, active: tuple[str, ...]) -> dict[str, object]:
    layout_id = "ZERES_GENIII_1080_V1"
    values: dict[str, object] = {
        "CoreAMode": "ON" if "A" in active else "OFF",
        "CoreBMode": "ON" if "B" in active else "OFF",
        "CoreCMode": "ON" if "C" in active else "OFF",
        "CoreDMode": "ON" if "D" in active else "OFF",
    }

    for row in mapping.core_fields("HAITIAN_ZE_III", layout_id):
        if str(row.get("STATUS", "")).strip().upper() != "MAPPED":
            continue
        key = str(row.get("DATA KEY", "")).strip()
        if not key:
            continue
        if key.endswith(("InMoldType", "OutMoldType")):
            values[key] = "POS MOLD"
        elif key.endswith("OutDuringMold"):
            values[key] = "OFF"
        elif key.endswith(("InMode", "OutMode")):
            values[key] = "TIEMPO"
        elif key.endswith(("InStandby", "OutStandby")):
            values[key] = "ON"
        else:
            values[key] = 1
    return values


def test_310b_three_cores_fill_left_center_right_slots():
    mapping = MappingCatalog(PROJECT_ROOT / "data" / "Mapeo.xlsx")
    builder = PlanBuilder(mapping)
    values = _core_values(mapping, ("A", "B", "C"))

    ops, active = builder._build_core_ops(
        "HAITIAN_ZE_III",
        "ZERES_GENIII_1080_V1",
        values,
        "H15",
        "M15",
    )

    assert active == ["A", "B", "C"]

    # Representative pressure cells for the three physical Noyo slots.
    assert any(op.kind == "WRITE" and op.cell == "O84" and op.value == 1 for op in ops)
    assert any(op.kind == "WRITE" and op.cell == "AH84" and op.value == 1 for op in ops)
    assert any(op.kind == "WRITE" and op.cell == "BA84" and op.value == 1 for op in ops)

    # The ON/OFF cells beside Noyo A/B/C must reflect the actual physical
    # slots, not retain the static ON values from the template.
    assert any(op.kind == "WRITE" and op.cell == "E78" and op.value == "ON" for op in ops)
    assert any(op.kind == "WRITE" and op.cell == "X78" and op.value == "ON" for op in ops)
    assert any(op.kind == "WRITE" and op.cell == "AQ78" and op.value == "ON" for op in ops)


def test_310b_one_active_core_sets_other_noyo_status_cells_off():
    mapping = MappingCatalog(PROJECT_ROOT / "data" / "Mapeo.xlsx")
    builder = PlanBuilder(mapping)
    values = _core_values(mapping, ("A",))

    ops, active = builder._build_core_ops(
        "HAITIAN_ZE_III",
        "ZERES_GENIII_1080_V1",
        values,
        "H15",
        "M15",
    )

    assert active == ["A"]
    status = {op.cell: op.value for op in ops if op.kind == "WRITE" and op.cell in {"E78", "X78", "AQ78"}}
    assert status == {"E78": "ON", "X78": "OFF", "AQ78": "OFF"}


def test_geniii_500_single_noyo_status_is_dynamic():
    mapping = MappingCatalog(PROJECT_ROOT / "data" / "Mapeo.xlsx")
    builder = PlanBuilder(mapping)
    layout_id = "ZERES_GENIII_V1"

    values = {
        "CoreAMode": "ON",
        "CoreBMode": "OFF",
        "CoreCMode": "OFF",
        "CoreDMode": "OFF",
    }
    for row in mapping.core_fields("HAITIAN_ZE_III", layout_id):
        if str(row.get("STATUS", "")).strip().upper() != "MAPPED":
            continue
        key = str(row.get("DATA KEY", "")).strip()
        if not key:
            continue
        if key.endswith(("InMoldType", "OutMoldType")):
            values[key] = "POS MOLD"
        elif key.endswith("OutDuringMold"):
            values[key] = "OFF"
        elif key.endswith(("InMode", "OutMode")):
            values[key] = "TIEMPO"
        elif key.endswith(("InStandby", "OutStandby")):
            values[key] = "ON"
        else:
            values[key] = 1

    ops_on, active_on = builder._build_core_ops(
        "HAITIAN_ZE_III", layout_id, values, "H16", "M16"
    )
    assert active_on == ["A"]
    assert any(op.kind == "WRITE" and op.cell == "E52" and op.value == "ON" for op in ops_on)

    values.update({"CoreAMode": "OFF"})
    ops_off, active_off = builder._build_core_ops(
        "HAITIAN_ZE_III", layout_id, values, "H16", "M16"
    )
    assert active_off == []
    assert any(op.kind == "WRITE" and op.cell == "E52" and op.value == "OFF" for op in ops_off)


def test_geniii_800_has_dynamic_noyo_status_cells_configured():
    mapping = MappingCatalog(PROJECT_ROOT / "data" / "Mapeo.xlsx")
    settings = mapping.core_settings("HAITIAN_ZE_III", "ZERES_GENIII_800_V1")

    assert settings["LEFT_STATUS_CELL"] == "E78"
    assert settings["CENTER_STATUS_CELL"] == "X78"
    assert settings["RIGHT_STATUS_CELL"] == "AQ78"


def test_310b_four_active_cores_are_blocked():
    mapping = MappingCatalog(PROJECT_ROOT / "data" / "Mapeo.xlsx")
    builder = PlanBuilder(mapping)
    values = _core_values(mapping, ("A", "B", "C", "D"))

    with pytest.raises(ProcessWarning, match="maximo 3"):
        builder._build_core_ops(
            "HAITIAN_ZE_III",
            "ZERES_GENIII_1080_V1",
            values,
            "H15",
            "M15",
        )
