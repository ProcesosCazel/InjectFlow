from pathlib import Path

from app.catalogs import MappingCatalog
from app.plan import PlanBuilder


def _mapping() -> MappingCatalog:
    fixtures = Path(__file__).parent / "fixtures"
    return MappingCatalog(fixtures / "Mapeo_schema_fixture.xlsx")


def _values_for_plan(mapping: MappingCatalog) -> dict[str, object]:
    values: dict[str, object] = {r.key: 1 for r in mapping.mapping_rows if r.key}
    values.update({
        "MoldCloseZones": 6,
        "MoldOpenZones": 6,
        "InjectionZones": 6,
        "HoldingZones": 4,
        "ChargeZones": 3,
        "EjectorZonesFWD": 3,
        "EjectorZonesBWD": 3,
    })
    for n in range(1, 41):
        values[f"HRSZone{n}Mode"] = "ON"
    for n in range(1, 9):
        values[f"BarrelZone{n}Mode"] = "ON"
    values.update({"CoreAMode": "ON", "CoreBMode": "OFF", "CoreCMode": "OFF", "CoreDMode": "OFF"})
    for row in mapping.core_field_rows:
        key = row.get("DATA KEY")
        if key:
            values.setdefault(str(key), 1)
    return values


def test_required_keys_exclude_resul_in_param_only_mode():
    mapping = _mapping()
    all_keys = mapping.required_keys_for("NAVE_1", "NAVE1_V1")
    param_only_keys = mapping.required_keys_for(
        "NAVE_1", "NAVE1_V1", enabled_sources={"PARAM", "MANUAL"}
    )

    assert "CycleTime" in all_keys
    assert "ChargeSafetyTime" in all_keys
    assert "CycleTime" not in param_only_keys
    assert "ChargeSafetyTime" not in param_only_keys
    assert "MoldCloseZones" in param_only_keys
    assert "CoreAMode" in param_only_keys


def test_plan_excludes_resul_destination_cells_in_param_only_mode():
    mapping = _mapping()
    values = _values_for_plan(mapping)
    plan = PlanBuilder(mapping).build(
        "NAVE_1",
        "NAVE1_V1",
        values,
        {"MoldNumber": "TEST", "MachineNumber": "125"},
        enabled_sources={"PARAM", "MANUAL"},
    )
    cells = {op.cell for op in plan.operations}

    for cell in {"AG69", "AG73", "AG74", "AG75", "AG76", "AG77", "AP41"}:
        assert cell not in cells


def test_param_only_clears_formula_depending_on_resultant(tmp_path: Path):
    from openpyxl import Workbook
    from app.main import _param_only_formula_cleanup

    mapping = _mapping()
    template = tmp_path / "template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "TEST"
    ws["AG71"] = "=(3600/AG69)"
    ws["B1"] = "=1+1"
    wb.save(template)

    ops = _param_only_formula_cleanup(
        template,
        "TEST",
        mapping,
        "NAVE_1",
        "NAVE1_V1",
    )
    assert any(op.kind == "CLEAR" and op.cell == "AG71" for op in ops)
    assert not any(op.cell == "B1" for op in ops)
