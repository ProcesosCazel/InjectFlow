"""Jupiter integration contract. Synthetic inputs are not HMI validation."""
from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
import copy
import xml.etree.ElementTree as ET

import openpyxl
import pytest

import app.main as engine
from app.catalogs import DataCatalog, MappingCatalog
from app.errors import ConfigurationError, ProcessWarning, SourceDataError
from app.input_policy import JUPITER_GROUP, policy_for_group
from app.jupiter_xml_parser import JupiterXmlParser
from app.plan import PlanBuilder
from app.resolver import DataResolver
from web_api import InjectFlowAPI
import web_api

ROOT = Path(__file__).resolve().parents[1]
LAYOUT = 'JUPITER_TWOSHOT_1080_V1'
SOURCES = {'PARAM', 'MOLD', 'MANUAL'}


def write_xml(path, variables=None):
    root = ET.Element('HMI_Data', {'Version': '1.2'})
    ET.SubElement(root, 'Comment', {'Text': '1689 test'})
    group = ET.SubElement(root, 'VarGroup', {'Name': 'VG_MoldData'})
    for name, value in (variables or {'system.sv_iCavities': 4}).items():
        item = ET.SubElement(group, 'Variable', {'Name': name})
        ET.SubElement(item, 'Value').text = str(value).lower() if isinstance(value, bool) else str(value)
    ET.ElementTree(root).write(path, encoding='utf-8', xml_declaration=True)
    return path


@pytest.fixture(scope='module')
def catalogs():
    return DataCatalog(ROOT/'data/Data.xlsx'), MappingCatalog(ROOT/'data/Mapeo.xlsx')


@pytest.fixture
def xml_path(tmp_path):
    return write_xml(tmp_path/'mold.XML', {
        'Mold1.sv_MoldFwdProfVisSrc.Profile.iNoOfPoints': 5,
        'Mold1.sv_MoldBwdProfVis.Profile.iNoOfPoints': 1,
        'Injection1.sv_InjectProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_InjectProfVis.Profile.iNoOfPoints': 1,
        'Injection1.sv_HoldProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_HoldProfVis.Profile.iNoOfPoints': 1,
        'HeatingMold1.sv_ZoneRetain1.bUsed': True,
        'HeatingMold1.sv_ZoneRetain1.rSetValVis': 260,
        'Core2.sv_CoreMode.CoreType': 1,
        'Core2.sv_CoreMode.CoreControlIn': 1,
        'Core2.sv_CoreMode.CoreControlOut': 1,
        'CentralCoordination1.sv_CoreData[2].InMode': 0,
        'CentralCoordination1.sv_CoreData[2].OutMode': 2,
        'Core2.sv_CoreOutput.NormalIn.Pressure.Output.rOutputValue': 37,
        'ValveGate1.sv_ValveGateData.ValveGateDataArray[9].bUsed': True,
        'ValveGate1.sv_ValveGateData.ValveGateDataArray[13].bUsed': True,
    })


@pytest.fixture
def values(catalogs, xml_path):
    data, mapping = catalogs
    parser = JupiterXmlParser()
    resolver = DataResolver(data)
    return resolver.resolve_param(parser.parse(xml_path),
        required_keys=mapping.required_keys_for(JUPITER_GROUP, LAYOUT, enabled_sources=SOURCES))


def build(catalogs, values):
    data, mapping = catalogs
    return PlanBuilder(mapping, data).build(JUPITER_GROUP, LAYOUT, values,
        {'MachineNumber': '308B', 'MoldNumber': 'I-1689'}, enabled_sources=SOURCES)


@pytest.mark.parametrize('group,ext,results,control', [
    ('HAITIAN_ZE_V', '.dat', True, True),
    ('HAITIAN_ZE_III', '.dat', True, True),
    (JUPITER_GROUP, '.xml', False, False),
])
def test_family_input_policy(group, ext, results, control):
    policy = policy_for_group(group)
    assert (policy.param_extension, policy.supports_resultants, policy.requires_injection_control) == (ext, results, control)
    policy.validate_param_path(Path('sample'+ext.upper()))
    with pytest.raises(SourceDataError):
        policy.validate_param_path(Path('sample.csv'))


def test_core_b_stays_right_even_when_core_a_is_off(catalogs, values):
    plan = build(catalogs, values)
    assert plan.active_cores == ('B',)
    layout = catalogs[1].core_layout(JUPITER_GROUP, LAYOUT)
    target = next(r['CELL'] for r in layout if r['SLOT']=='RIGHT' and r['FIELD']=='PRESSURE' and r['DIRECTION']=='IN')
    assert any(o.kind=='WRITE' and o.cell==target and o.value==Decimal(37) for o in plan.operations)
    left = {r['CELL'] for r in layout if r['SLOT']=='LEFT'}
    assert not any(o.kind=='WRITE' and o.cell in left for o in plan.operations)


def test_jupiter_fixed_end_rows_are_not_compacted_by_stage_counts(catalogs, values):
    plan = build(catalogs, values)
    mapping = catalogs[1]
    rows = {r.key: r for r in mapping.rows_for(JUPITER_GROUP, LAYOUT)}

    fixed_end_keys = {
        'MoldOpenEndPressure',
        'Injection1EndPressure', 'Injection2EndPressure',
        'Hold1EndPressure', 'Hold2EndPressure',
        'Charge1EndRPM', 'Charge2EndRPM',
    }
    for key in fixed_end_keys:
        cell = rows[key].cell
        assert any(o.kind == 'WRITE' and o.cell == cell for o in plan.operations), key

    # The fixture declares one profile point for open/injection/hold. That one
    # point is End, so numbered Stage 1 rows must remain inactive rather than
    # receiving the End value by compaction.
    numbered_keys = {
        'MoldOpenStage1Pressure',
        'Injection1Stage1Pressure', 'Injection2Stage1Pressure',
        'Hold1Stage1Pressure', 'Hold2Stage1Pressure',
    }
    for key in numbered_keys:
        cell = rows[key].cell
        assert not any(o.kind == 'WRITE' and o.cell == cell for o in plan.operations), key


def test_direct_valve_numbering_includes_9_and_does_not_shift_13(catalogs, values):
    plan = build(catalogs, values)
    assert plan.active_valve_gates == (9, 13)
    gate_cells = {r['CELL'] for r in catalogs[1].valve_gate_layout(JUPITER_GROUP, LAYOUT) if r['FIELD']=='GATE_NO'}
    assert [o.value for o in plan.operations if o.kind=='WRITE' and o.cell in gate_cells] == [9, 13]


def test_jupiter_modes_use_jupiter_position_time_map_not_zeres(catalogs, values):
    # Feed raw controller values directly to the planner to verify that its
    # Jupiter branch is independent from the Zeres generation StateMaps.
    values['ValveGate9StartCon'] = Decimal(1)
    values['ValveGate9StopCon'] = Decimal(0)
    plan = build(catalogs, values)
    layout = catalogs[1].valve_gate_layout(JUPITER_GROUP, LAYOUT)
    start_cell = next(r['CELL'] for r in layout
                      if r['SLOT']=='SLOT1' and r['FIELD']=='START_CON')
    stop_cell = next(r['CELL'] for r in layout
                     if r['SLOT']=='SLOT1' and r['FIELD']=='STOP_CON')
    assert any(o.cell==start_cell and o.kind=='WRITE' and o.value=='TIME' for o in plan.operations)
    assert any(o.cell==stop_cell and o.kind=='WRITE' and o.value=='POSITION' for o in plan.operations)


def test_jupiter_over_nine_active_gates_blocks(catalogs, values):
    for n in range(1,11): values[f'ValveGate{n}State']='ON'
    with pytest.raises(ProcessWarning): build(catalogs, values)


def test_plan_does_not_touch_resultants_or_fixed_tolerance_merges(catalogs, values):
    plan = build(catalogs, values)
    protected = {f'{c}{r}' for c in ('BC','BH','BM','BR','BW') for r in (52,54,57,59)}
    protected |= {'AC65','AC66','AC67'} | {f'{c}{r}' for c in ('AC','AG') for r in range(70,75)}
    wb = openpyxl.load_workbook(ROOT/'plantillas/Haitian Jupiter TwoShot_1080.xlsx')
    try:
        ws=wb.active
        lookup={cell.coordinate:ws.cell(area.min_row,area.min_col).coordinate
                for area in ws.merged_cells.ranges
                for row in ws.iter_rows(min_row=area.min_row,max_row=area.max_row,
                                       min_col=area.min_col,max_col=area.max_col) for cell in row}
        assert not any(lookup.get(o.cell,o.cell) in protected for o in plan.operations)
        assert all(lookup.get(o.cell,o.cell)==o.cell for o in plan.operations)
    finally: wb.close()


def test_cool_prevent_time_is_template_owned_and_not_mapped(catalogs, values):
    assert 'Barrel1CoolPreventTime' not in values and 'Barrel2CoolPreventTime' not in values
    targets={r.cell for r in catalogs[1].rows_for(JUPITER_GROUP,LAYOUT)
             if r.key in {'Barrel1CoolPreventTime','Barrel2CoolPreventTime'}}
    assert targets == set()
    plan = build(catalogs, values)
    assert not any(o.cell in {'BI55','BI61'} and o.kind in {'WRITE','CLEAR'} for o in plan.operations)


@pytest.mark.parametrize('bad', ['NaN', 'sNaN', 'Infinity', '-Infinity'])
def test_parser_rejects_nonfinite_setpoints(tmp_path, bad):
    path=write_xml(tmp_path/'bad.xml', {'HeatingNozzle1.sv_ZoneRetain1.rSetValVis':bad})
    with pytest.raises(SourceDataError,match='Invalid numeric'): JupiterXmlParser().parse(path)


@pytest.mark.parametrize('bad', ['1.5','NaN','Infinity'])
def test_parser_rejects_noninteger_stage_count(tmp_path, bad):
    path=write_xml(tmp_path/'bad.xml', {'Injection1.sv_InjectProfVis.Profile.iNoOfPoints':bad})
    with pytest.raises(SourceDataError,match='Invalid integer'): JupiterXmlParser().parse(path)


def args_for(tmp_path, xml_path, machine='308B'):
    return SimpleNamespace(machine=machine,mold='I-1689',injection_control_mode=None,
        data=ROOT/'data/Data.xlsx',mapping=ROOT/'data/Mapeo.xlsx',molds=ROOT/'data/Moldes.xlsx',
        param=xml_path,resul=tmp_path/'missing.csv',param_only=False,
        templates_dir=ROOT/'plantillas',output=tmp_path/'sheet.xlsx',history=tmp_path/'history.csv',
        report=tmp_path/'report.txt',plan_csv=tmp_path/'plan.csv',plan_only=True,
        allow_inactive_machine=False,allow_template_mismatch=False)


@pytest.mark.parametrize('machine', ['308B','309'])
def test_engine_routes_xml_forces_param_only_without_manual_control(tmp_path, xml_path, monkeypatch, catalogs, machine, capsys):
    data,mapping=catalogs
    monkeypatch.setattr(engine,'DataCatalog',lambda path:data)
    monkeypatch.setattr(engine,'MappingCatalog',lambda path:mapping)
    def forbidden(*args,**kwargs): raise AssertionError('DAT/CSV/prompt/writer must not run')
    monkeypatch.setattr(engine,'ParamDatParser',forbidden)
    monkeypatch.setattr(engine,'ResulCsvParser',forbidden)
    monkeypatch.setattr(engine,'_prompt_injection_control_mode',forbidden)
    monkeypatch.setattr(engine,'ExcelComWriter',forbidden)
    args=args_for(tmp_path,xml_path,machine)
    assert engine.run(args)==0
    log=capsys.readouterr().out
    assert 'SOLO PARAMETROS XML' in log and 'Valve Gates activos: 9, 13' in log
    assert args.plan_csv.exists() and not args.output.exists() and not args.history.exists()
    assert 'TOTAL: 1 ZONAS' in args.plan_csv.read_text(encoding='utf-8-sig')
    assert machine in mapping.machines


def test_engine_passes_machine_constants_to_jupiter_charge_conversion(tmp_path, catalogs):
    data, mapping = catalogs
    xml = write_xml(tmp_path/'charge.xml', {
        'Mold1.sv_MoldFwdProfVisSrc.Profile.iNoOfPoints': 3,
        'Mold1.sv_MoldBwdProfVis.Profile.iNoOfPoints': 1,
        'Injection1.sv_InjectProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_InjectProfVis.Profile.iNoOfPoints': 1,
        'Injection1.sv_HoldProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_HoldProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_PlastProfVis.Profile.iNoOfPoints': 1,
        'Injection2.sv_PlastProfVis.Profile.Points[1].rRotation': 57.595867,
        'Injection2.sv_PlastProfVis.Profile.Points[1].rPressure': 90,
        'Injection2.sv_PlastProfVis.Profile.Points[1].rBackPressure': 5,
        'Injection2.sv_PlastProfVis.Profile.Points[2].rStartPos': 186.53206,
    })
    args = args_for(tmp_path, xml, '308B')
    assert engine.run(args) == 0

    rpm_cell = next(
        r.cell for r in mapping.rows_for(JUPITER_GROUP, LAYOUT)
        if r.key == 'Charge2EndRPM'
    )
    plan_rows = args.plan_csv.read_text(encoding='utf-8-sig').splitlines()
    assert any(rpm_cell in row and '100' in row for row in plan_rows)


def test_development_override_never_permits_excel_write(tmp_path,xml_path):
    args=args_for(tmp_path,xml_path);args.plan_only=False;args.allow_inactive_machine=True
    with pytest.raises(ValueError,match='solo se permite'):engine.run(args)


def test_jupiter_machine_is_active_without_override(catalogs):
    m=catalogs[1]
    group, template = m.resolve_machine('308B')
    assert group == JUPITER_GROUP
    assert template.active
    assert m.machines['308B'] == JUPITER_GROUP
    assert m.machines['309'] == JUPITER_GROUP


def test_engine_accepts_user_mold_even_when_xml_comment_differs(tmp_path,xml_path,capsys):
    args=args_for(tmp_path,xml_path);args.mold='I-1831'
    assert engine.run(args) == 0
    log = capsys.readouterr().out
    assert 'no coincide' not in log
    assert 'no encontrado en Moldes.xlsx' not in log
    assert 'Molde: I-1831' in log


@pytest.fixture
def api(tmp_path,monkeypatch,catalogs):
    # Jupiter is active in the real-test catalog; use that catalog through the bridge.
    m=catalogs[1]
    monkeypatch.setattr(web_api,'MappingCatalog',lambda path:m)
    (tmp_path/'mapping.xlsx').touch()
    return InjectFlowAPI(project_root=tmp_path,mapping_path=tmp_path/'mapping.xlsx')


def test_bridge_exposes_family_capabilities_without_changing_default_catalog(api):
    options=api.get_manual_selectors()['machines']
    j=next(o for o in options if o['value']=='308B')
    assert j['inputMode']=='XML_ONLY' and j['paramExtension']=='.xml'
    assert j['supportsResultants'] is False and j['requiresInjectionControl'] is False
    z=next(o for o in options if o['value']=='112C')
    assert z['paramExtension']=='.dat' and z['supportsResultants'] is True


def test_changing_family_clears_stale_paths_not_files(api,tmp_path,xml_path):
    dat=tmp_path/'param.dat';dat.write_text('sample')
    csv=tmp_path/'result.csv';csv.write_text('sample')
    assert api._accept_input_path('param',dat,source='test')['ok']
    assert api._accept_input_path('resul',csv,source='test')['ok']
    change=api.set_machine('308B')
    assert change['ok'] and set(change['clearedInputs'])=={'param','resul'}
    assert dat.exists() and csv.exists()
    assert api._accept_input_path('param',xml_path,source='drop')['ok']
    assert api._accept_input_path('resul',csv,source='drop')['error']['code']=='resultants_not_supported'
    assert api._accept_input_path('param',dat,source='dialog')['error']['code']=='wrong_extension'
    change=api.set_machine('112C')
    assert change['files']['param'] is None and change['policy']['supportsResultants']
    assert api._accept_input_path('resul',csv,source='drop')['ok']


def test_bridge_generation_forces_xml_only_despite_stale_csv_and_client_flags(api,xml_path,monkeypatch,tmp_path):
    assert api.set_machine('308B')['ok']
    assert api._accept_input_path('param',xml_path,source='test')['ok']
    api._input_files['resul']={'path':str(tmp_path/'nonexistent.csv'),'name':'nonexistent.csv'}
    captured=[]
    def fake_run(args):
        captured.append(args);args.output.write_bytes(b'test-writer-only')
        print('Operaciones planeadas: 123\nVerificacion Excel OK: 100 celdas finales')
        return 0
    monkeypatch.setattr(engine,'run',fake_run)
    result=api.generate_sheet({'machine':'308B','mold':'I-1689','paramOnly':False,'allow_inactive_machine':True})
    assert result['ok'] and result['paramOnly'] and result['engineMode']=='SOLO PARAMETROS XML'
    assert captured[0].resul is None and captured[0].param_only is True
    assert captured[0].injection_control_mode is None
    assert result['injectionControl']['code'] is None
    assert not hasattr(captured[0],'allow_inactive_machine')


def test_bridge_zeres_still_requires_manual_control(api,xml_path,tmp_path):
    dat=tmp_path/'param.dat';dat.write_text('sample')
    assert api.set_machine('112C')['ok']
    assert api._accept_input_path('param',dat,source='test')['ok']
    result=api.generate_sheet({'machine':'112C','mold':'I-1689'})
    assert not result['ok'] and result['error']['code']=='missing_inputs'


def test_bridge_exposes_active_jupiter_machines(tmp_path,catalogs):
    bridge=InjectFlowAPI(project_root=tmp_path,mapping_path=ROOT/'data/Mapeo.xlsx')
    assert bridge.set_machine('308B')['ok']
    values={r['value'] for r in bridge.get_manual_selectors()['machines']}
    assert {'308B','309'} <= values


def test_bridge_rejects_input_mutations_during_generation(api,xml_path):
    api._generation_running=True
    assert not api.set_machine('112C')['ok']
    assert api.clear_input_file('param')['error']['code']=='generation_busy'
    assert api._accept_input_path('param',xml_path,source='drop')['error']['code']=='generation_busy'


def test_native_dialog_filter_is_xml_and_incompatible_dialog_result_is_rejected(api,xml_path,monkeypatch):
    import sys
    fake_webview=SimpleNamespace(FileDialog=SimpleNamespace(OPEN=0))
    monkeypatch.setitem(sys.modules,'webview',fake_webview)
    calls=[]
    class Window:
        def create_file_dialog(self,**kwargs):
            calls.append(kwargs)
            api.set_machine('112C')  # simulate change while the dialog is open
            return [str(xml_path)]
    api._attach_window(Window());api.set_machine('308B')
    result=api.select_input_file('param')
    assert '*.xml' in calls[0]['file_types'][0]
    assert result['error']['code']=='machine_changed'
    assert api.get_input_files()['files']['param'] is None


def test_cli_wrong_extension_blocks_before_parser(tmp_path, xml_path, monkeypatch, catalogs):
    args=args_for(tmp_path,xml_path);args.param=tmp_path/'mistaken.dat'
    args.param.write_text('not xml')
    def forbidden(*args,**kwargs): raise AssertionError('Parser must not run')
    monkeypatch.setattr(engine,'JupiterXmlParser',forbidden)
    with pytest.raises(SourceDataError,match='archivo .xml'):engine.run(args)


def test_jupiter_incomplete_profile_does_not_generate_a_successful_plan(catalogs,values):
    # No stage-zero rule exists for mold closing in the official template.
    # Incomplete controller data must not silently create an all-zero process.
    values['MoldCloseStages']=Decimal(0)
    with pytest.raises(ConfigurationError,match='MoldCloseStages'):build(catalogs,values)


@pytest.fixture
def browser_page(api, xml_path, tmp_path, monkeypatch):
    """Optional headless Chromium test; native Excel/WebView2 is NOT simulated."""
    import json
    import os
    import shutil
    import re
    import base64
    playwright=pytest.importorskip('playwright.sync_api')
    chromium=shutil.which('chromium') or shutil.which('chromium-browser')
    dat=tmp_path/'params.dat';dat.write_text('test only')
    csv=tmp_path/'results.csv';csv.write_text('test only')
    class Window:
        def create_file_dialog(self,**kwargs):
            filt=kwargs['file_types'][0]
            return [str(csv if '*.csv' in filt else xml_path if '*.xml' in filt else dat)]
    import sys
    monkeypatch.setitem(sys.modules,'webview',SimpleNamespace(FileDialog=SimpleNamespace(OPEN=0)))
    api._attach_window(Window())
    names=['bootstrap','get_manual_selectors','notify_ui_ready','get_input_files',
           'set_machine','ping','get_latest_output','get_status','clear_input_file',
           'select_input_file']
    with playwright.sync_playwright() as pw:
        options={'headless':True,'args':['--no-sandbox']}
        if chromium:options['executable_path']=chromium
        try: browser=pw.chromium.launch(**options)
        except playwright.Error:pytest.skip('Chromium for UI tests is not installed')
        page=browser.new_page(viewport={'width':1365,'height':1050})
        page.set_default_timeout(7000)
        page.on('pageerror',lambda exc:errors.append(str(exc)))
        errors=[]
        for name in names:page.expose_function('test_'+name,getattr(api,name))
        # Offline harness: use local asset bytes; no browser network access,
        # native file-dialog access, or production machine activation.
        html=(ROOT/'web/index.html').read_text(encoding='utf-8')
        html=re.sub(r'<script[^>]*src=[^>]*></script>', '', html)
        css=(ROOT/'web/css/styles.css').read_text(encoding='utf-8')
        html=re.sub(r'<link[^>]*href=[^>]*styles\.css[^>]*>',lambda _: '<style>'+css+'</style>',html)
        logo=base64.b64encode((ROOT/'web/assets/LogoCazel.webp').read_bytes()).decode()
        html=re.sub(r'(?:\./)?assets/LogoCazel\.webp','data:image/webp;base64,'+logo,html)
        page.route('**/*',lambda route:route.abort())
        page.set_content(html)
        page.evaluate('window.pywebview={api:Object.fromEntries('+json.dumps(names)+
                      '.map(n=>[n,(...args)=>window["test_"+n](...args)]))};')
        page.add_script_tag(content=(ROOT/'web/js/main.js').read_text(encoding='utf-8'))
        page.evaluate('InjectFlowUI.init()')
        page.wait_for_function('window.InjectFlowUI?.getState().backend.manualSelectorsLoaded')
        yield page,api,errors
        if os.environ.get('INJECTFLOW_TEST_ARTIFACTS'):
            directory=Path(os.environ['INJECTFLOW_TEST_ARTIFACTS']);directory.mkdir(parents=True,exist_ok=True)
            page.screenshot(path=str(directory/(tmp_path.name+'.png')),full_page=True)
        browser.close()


@pytest.mark.parametrize("width", [1365,800])
def test_ui_jupiter_hides_csv_and_manual_control(browser_page,width):
    page,api,errors=browser_page
    page.set_viewport_size({"width":width,"height":1050})
    page.select_option('#machine-select','308B')
    page.wait_for_function('!InjectFlowUI.getState().machineSyncPending && InjectFlowUI.getState().machine==="308B"')
    assert not page.locator('#resul-upload-panel').is_visible()
    assert not page.locator('#injection-control-panel').is_visible()
    assert page.locator('#resul-upload-button').is_disabled()
    assert 'XML' in page.locator('#param-upload-title').inner_text()
    page.click('#param-upload-button')
    page.wait_for_function('!!InjectFlowUI.getState().paramFile')
    page.fill('#mold-input','I-1689')
    assert page.locator('#generate-button').is_enabled()
    summary=page.locator('#summary-content').inner_text()
    assert 'XML' in summary and 'Resul.csv' not in summary
    assert api._input_policy.supports_resultants is False
    assert page.locator('#process-data-title').inner_text().startswith('2.')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert not errors


def test_ui_switching_back_to_zeres_restores_csv_and_control(browser_page):
    page,api,errors=browser_page
    page.evaluate('async()=>await InjectFlowUI.setMachine("308B")')
    page.click('#param-upload-button')
    page.wait_for_function('!!InjectFlowUI.getState().paramFile')
    page.evaluate('async()=>await InjectFlowUI.setMachine("112C")')
    assert page.locator('#resul-upload-panel').is_visible()
    assert page.locator('#injection-control-panel').is_visible()
    assert page.evaluate('InjectFlowUI.getState().paramFile') is None
    page.click('#param-upload-button')
    page.wait_for_function('!!InjectFlowUI.getState().paramFile')
    page.fill('#mold-input','I-1689')
    assert page.locator('#generate-button').is_disabled()
    page.select_option('#inj-ctrl-select',label='Modo velocidad')
    assert page.locator('#generate-button').is_enabled()
    assert api._input_policy.supports_resultants is True
    assert not errors


def test_ui_fast_machine_changes_leave_backend_and_frontend_consistent(browser_page):
    page,api,errors=browser_page
    page.evaluate('async()=>await Promise.all([InjectFlowUI.setMachine("308B"),InjectFlowUI.setMachine("112C"),InjectFlowUI.setMachine("309")])')
    state=page.evaluate('InjectFlowUI.getState()')
    assert state['machine']=='309' and api._selected_machine=='309'
    assert state['inputPolicy']['paramExtension']=='.xml'
    assert not state['machineSyncPending'] and not state['machineSyncFailed']
    assert not errors


def test_jupiter_transfer_switches_render_checks_and_blank_disabled_thresholds(catalogs, values):
    values.update({
        'Injection1UsePosition': 'ON',
        'Injection1UseTimer': 'OFF',
        'Injection1UsePressure': 'OFF',
        'Injection1ScrewPosition': Decimal(45),
        'Injection1TimeSet': Decimal(6),
        'Injection1PressureSet': Decimal(100),
    })
    plan = build(catalogs, values)
    writes = {(o.cell, o.value) for o in plan.operations if o.kind == 'WRITE'}
    assert ('B31', '✓') in writes
    assert ('B32', None) in writes and ('B33', None) in writes
    assert not any(o.kind == 'WRITE' and o.cell in {'L32', 'L33'} for o in plan.operations)
    assert ('L31', Decimal(45)) in writes


def test_jupiter_valve_gate_boolean_fields_render_as_checkmarks(catalogs, values):
    values['ValveGate9StartUse'] = 'ON'
    values['ValveGate9OnWhenHoldingState'] = 'ON'
    plan = build(catalogs, values)
    assert any(o.kind == 'WRITE' and o.value == '✓' and o.note == 'VALVE_GATE_9:ValveGate9StartUse'
               for o in plan.operations)
    assert any(o.kind == 'WRITE' and o.value == '✓' and o.note == 'VALVE_GATE_9:ValveGate9OnWhenHoldingState'
               for o in plan.operations)


def test_jupiter_mold_protect_and_hold_end_become_active_when_written(catalogs, values):
    values.update({
        'MoldProtectStagePressure': Decimal(105),
        'MoldProtectStageVelocity': Decimal(28),
        'MoldProtectStagePosition': Decimal(0),
        'Holding1Stages': Decimal(1),
        'Hold1EndPressure': Decimal(25),
        'Hold1EndVelocity': Decimal(10),
        'Hold1EndTime': Decimal(3),
        'Holding2Stages': Decimal(1),
        'Hold2EndPressure': Decimal(30),
        'Hold2EndVelocity': Decimal(10),
        'Hold2EndTime': Decimal(3),
    })
    plan = build(catalogs, values)
    for cell in {'G20','M20','R20','BD29','BJ29','BO29','CB29','CH29','CM29'}:
        assert any(o.kind == 'COPY_FILL' and o.cell == cell and o.ref_cell == 'G17' for o in plan.operations), cell
        assert any(o.kind == 'WRITE' and o.cell == cell for o in plan.operations), cell


def test_jupiter_fixed_profile_labels_live_in_locked_template(catalogs, values):
    plan = build(catalogs, values)
    fixed_cells = {
        'B26':1,'B27':2,'B28':3,'B29':'End',
        'Z26':1,'Z27':2,'Z28':3,'Z29':'End',
        'BA26':1,'BA27':2,'BA28':3,'BA29':'End',
        'BY26':1,'BY27':2,'BY28':3,'BY29':'End',
    }
    # The template is authoritative for static labels. The plan must not rewrite them.
    assert not any(o.cell in fixed_cells for o in plan.operations)

    wb = openpyxl.load_workbook(ROOT/'plantillas/Haitian Jupiter TwoShot_1080.xlsx', data_only=False)
    try:
        ws = wb.active
        for cell, expected in fixed_cells.items():
            assert ws[cell].value == expected, cell
    finally:
        wb.close()


def test_jupiter_template_uses_locked_final_fill_refs(catalogs):
    wb = openpyxl.load_workbook(ROOT/'plantillas/Haitian Jupiter TwoShot_1080.xlsx')
    try:
        ws = wb.active
        active = ws['G17'].fill
        inactive = ws['G20'].fill
        assert active.fill_type is None
        assert inactive.fill_type == 'solid'
        assert inactive.fgColor.type == 'rgb'
        assert inactive.fgColor.rgb == 'FF7F7F7F'
    finally:
        wb.close()
    rules = [r for r in catalogs[1].zone_rules
             if r.template_group == JUPITER_GROUP and r.layout_id == LAYOUT]
    assert rules
    assert all(r.active_fill_ref == 'G17' and r.inactive_fill_ref == 'G20' for r in rules)
    settings = catalogs[1].core_settings(JUPITER_GROUP, LAYOUT)
    assert settings['ACTIVE_FILL_REF'] == 'G17'
    assert settings['INACTIVE_FILL_REF'] == 'G20'

