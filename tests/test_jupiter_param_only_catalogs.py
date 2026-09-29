"""Catalog contract for Jupiter XML-only real-test routing."""
from collections import Counter
from pathlib import Path

import openpyxl
import pytest

from app.catalogs import DataCatalog, MappingCatalog
from app.template_validation import validate_template

ROOT = Path(__file__).resolve().parents[1]
GROUP = 'HAITIAN_JUPITER_TWOSHOT'
LAYOUT = 'JUPITER_TWOSHOT_1080_V1'

@pytest.fixture(scope='module')
def catalogs():
    return DataCatalog(ROOT / 'data/Data.xlsx'), MappingCatalog(ROOT / 'data/Mapeo.xlsx')


def test_jupiter_csv_catalog_rules_are_removed(catalogs):
    data, _ = catalogs
    assert not any((r.source_name or '').startswith('Jupiter.') for r in data.resul_rules)
    for name in ('JUPITER_RESULT_EXTENSION', 'JUPITER_NORMALIZED_RESULT_PREFIX',
                 'JUPITER_RESULT_WINDOW', 'JUPITER_U2_TRANSFER_NORMALIZATION',
                 'JUPITER_U2_CUSHION_NORMALIZATION', 'JUPITER_U2_PRESSURE_NORMALIZATION'):
        assert name not in data.config
    assert data.cfg('JUPITER_PARAM_ONLY') is True
    assert data.cfg('JUPITER_PARAM_EXTENSION') == '.xml'
    assert data.cfg('RESUL_DEFAULT_WINDOW') == 20  # Zeres still supports resultants.


def test_jupiter_mapping_contains_no_resultants(catalogs):
    _, mapping = catalogs
    rows = mapping.rows_for(GROUP, LAYOUT)
    assert Counter(r.source for r in rows) == {'Param': 873, 'Mold': 6, 'Manual': 3}
    assert not any(r.source.upper() == 'RESUL' for r in rows)
    assert not any(r.cell in {'AC65', 'AC67', 'AC70', 'AG70', 'AC71', 'AG71',
                             'AC72', 'AG72', 'AC73', 'AG73', 'AC74', 'AG74'} for r in rows)


def test_jupiter_machines_are_xml_only_and_active_for_real_tests(catalogs):
    _, mapping = catalogs
    book = openpyxl.load_workbook(ROOT / 'data/Mapeo.xlsx', read_only=True, data_only=True)
    try:
        ws = book['Machines']
        headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        records = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
        matches = [r for r in records if str(r['MACHINE']) in {'308B', '309'}]
        assert len(matches) == 2
        for r in matches:
            assert r['INPUT MODE'] == 'XML_ONLY'
            assert str(r['ACTIVE']).upper() == 'TRUE'
        assert mapping.machines['308B'] == GROUP and mapping.machines['309'] == GROUP
        template = next(t for t in mapping.templates if t.layout_id == LAYOUT)
        assert template.active
    finally:
        book.close()


def test_jupiter_param_rule_coverage_is_retained(catalogs):
    data, mapping = catalogs
    normalized = [r for r in data.param_rules if (r.source_name or '').startswith('Jupiter.')]
    assert len(normalized) == 875
    assert len({r.rule_id for r in normalized}) == 875
    keys = {r.key for r in normalized}
    required = {r.key for r in mapping.rows_for(GROUP, LAYOUT) if r.source == 'Param'}
    assert required <= keys


def test_master_qad_cycle_is_not_a_measured_cycle(catalogs):
    data, mapping = catalogs
    master = {r.key:r.cell for r in mapping.rows_for(GROUP, LAYOUT) if r.source == 'Mold'}
    assert master == {'MoldDescription':'M3', 'MoldCustomer':'L4', 'MoldInternalCode':'N5',
                      'MoldCustomerPartNumber':'BR5', 'MoldCavities':'BY8', 'MoldQADCycleTime':'AM66'}
    assert any(r.key == 'CycleTime' for r in data.resul_rules)
    assert any(r.key == 'CycleTime' and r.source == 'Resul' and r.template_group != GROUP
               for r in mapping.mapping_rows)


def test_jupiter_template_anchors_and_empty_cycle_formula(catalogs):
    _, mapping = catalogs
    template = next(t for t in mapping.templates if t.layout_id == LAYOUT)
    path = ROOT / 'plantillas/Haitian Jupiter TwoShot_1080.xlsx'
    validation = validate_template(path, template, mapping, strict=True)
    assert validation.merged_ranges == 730
    assert not validation.warnings
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        ws = book.active
        assert ws['AC67'].value == '=IFERROR(3600/AC65,"")'
        assert ws['AC65'].value is None
        assert ws['AC66'].value == 2
    finally:
        book.close()
