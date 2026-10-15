"""Template revision 2026-10-15: layout checks, not machine/HMI validation.

The production Moldes.xlsx is intentionally not used as a fixture. The hashes
are read from Mapeo.xlsx; a new template revision must update its registration.
"""
from __future__ import annotations

from copy import copy
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import pytest

from app.catalogs import DataCatalog, MappingCatalog
from app.errors import TemplateError
from app.input_policy import JUPITER_GROUP
from app.main import _param_only_formula_cleanup
from app.plan import PlanBuilder
from app.template_validation import validate_template

ROOT = Path(__file__).resolve().parents[1]
LAYOUTS = (
    'NAVE1_V1', 'ZERES_GENIII_V1', 'ZERES_GENIII_800_V1',
    'ZERES_GENIII_1080_V1', 'JUPITER_TWOSHOT_1080_V1',
)
ADDRESS = re.compile(r'[A-Z]+[1-9][0-9]*')


@pytest.fixture(scope='module')
def catalogs():
    return DataCatalog(ROOT / 'data/Data.xlsx'), MappingCatalog(ROOT / 'data/Mapeo.xlsx')


@pytest.fixture(scope='module', params=LAYOUTS)
def template_case(request, catalogs):
    data, mapping = catalogs
    template = next(t for t in mapping.templates if t.layout_id == request.param)
    path = ROOT / 'plantillas' / template.template_file
    book = load_workbook(path)
    try:
        yield data, mapping, template, path, book.active
    finally:
        book.close()


def _anchors(ws):
    return {
        cell.coordinate: ws.cell(area.min_row, area.min_col).coordinate
        for area in ws.merged_cells.ranges
        for row in ws.iter_rows(min_row=area.min_row, max_row=area.max_row,
                                min_col=area.min_col, max_col=area.max_col)
        for cell in row
    }


def _refs(mapping, template):
    """Include static targets, dynamic layouts, labels and fill references."""
    result = set()
    fields = ('CELL', 'MOLD CELL', 'ACTIVE FILL REF', 'INACTIVE FILL REF', 'DESTINATION', 'VALUE')
    for table in mapping.tables.values():
        for row in table.rows:
            if row.get('TEMPLATE GROUP') != template.template_group or row.get('LAYOUT ID') != template.layout_id:
                continue
            for field in fields:
                for token in str(row.get(field) or '').replace('$', '').split('|'):
                    if ADDRESS.fullmatch(token):
                        result.add(token)
    return result


def _protected_notes(ws):
    return {
        cell.coordinate: cell.value
        for row in ws for cell in row
        if isinstance(cell.value, str)
        and ('NO SON AUDITABLES' in cell.value.upper()
             or cell.value == 'OBSERVACIONES'
             or cell.value.startswith('CICLO EN AUTOMATICO'))
    }


def _synthetic_values(mapping, template, active):
    """Boundary values for exercising the layout; NOT a production recipe."""
    group, layout = template.template_group, template.layout_id
    values = {key: 1 for key in mapping.required_keys_for(group, layout)}
    for key in list(values):
        if re.fullmatch(r'HRSZone\d+Mode', key):
            values[key] = 'ON' if active else 'OFF'
        elif key.endswith(('AfterActive', 'Enabled')) or re.fullmatch(r'Transfer[12].*Active', key):
            values[key] = 'ON' if active else 'OFF'
        elif key.startswith('Core') and key.endswith(('InMoldType', 'OutMoldType')):
            values[key] = 'POS MOLD' if active else 'CIER. BLOQ.'
        elif key.startswith('Core') and key.endswith(('InStandby', 'OutStandby', 'OutDuringMold')):
            values[key] = 'ON' if active else 'OFF'
        elif key.startswith('Core') and key.endswith(('InMode', 'OutMode')):
            values[key] = 'TIEMPO'
        elif re.fullmatch(r'ValveGate\d+.*State', key):
            values[key] = 'OFF'
        elif re.fullmatch(r'ValveGate\d+(StartCon|StopCon)', key):
            values[key] = 0
    controls = {}
    for rule in mapping.zone_rules:
        if rule.template_group == group and rule.layout_id == layout:
            controls.setdefault(rule.control_key, []).append(rule.control_value)
    for key, choices in controls.items():
        numeric = [v for v in choices if isinstance(v, (int, float))]
        if numeric:
            values[key] = max(numeric) if active else min(numeric)
    core = mapping.core_settings(group, layout)
    selectors = str(core.get('SELECTOR_KEYS', '')).split('|')
    limit = int(core['MAX_ACTIVE_CORES'])
    for index, key in enumerate(selectors):
        if key:
            values[key] = 'ON' if active and index < limit else 'OFF'
    gate = mapping.valve_gate_settings(group, layout)
    gate_numbers = sorted({int(row['GATE']) for row in mapping.valve_gate_fields(group, layout)})
    if group != JUPITER_GROUP:
        gate_numbers = [n for n in gate_numbers if n != 9]
    gate_limit = len({row['SLOT'] for row in mapping.valve_gate_layout(group, layout)})
    for number in gate_numbers[:gate_limit] if active else []:
        values[f'ValveGate{number}State'] = 'ON'
        values[f'ValveGate{number}OnWhenHoldingState'] = 'ON'
        values[f'ValveGate{number}StartUse'] = 'ON'
    values.update({'CycleTime': 60, 'MoldCustomer': 'TEST', 'MoldNumber': 'I-9000'})
    return values


def test_registered_template_passes_strict_validation(template_case):
    _, mapping, template, path, ws = template_case
    result = validate_template(path, template, mapping, strict=True)
    assert result.merged_ranges == len(ws.merged_cells.ranges)
    assert result.warnings == ()
    row = next(r for r in mapping.tables['Templates'].rows if r['LAYOUT ID'] == template.layout_id)
    assert ws.page_setup.orientation == row['ORIENTATION']
    assert ws.page_setup.scale == row['SCALE']


def test_registered_sha256_identifies_the_installed_template(template_case):
    _, _, template, path, _ = template_case
    assert template.reference_layout_signature
    assert sha256(path.read_bytes()).hexdigest() == template.reference_layout_signature


def test_static_dynamic_and_fill_references_are_valid_anchors(template_case):
    _, mapping, template, _, ws = template_case
    lookup = _anchors(ws)
    refs = _refs(mapping, template)
    assert refs
    min_col, min_row, max_col, max_row = range_boundaries(template.print_area)
    for address in refs:
        assert lookup.get(address, address) == address, address
        cell = ws[address]
        assert min_row <= cell.row <= max_row and min_col <= cell.column <= max_col, address
    builder = PlanBuilder(mapping)
    active_ref, inactive_ref = builder._generic_fill_refs(template.template_group, template.layout_id)
    assert copy(ws[active_ref].fill) != copy(ws[inactive_ref].fill)


def test_audit_notes_are_not_generation_destinations(template_case):
    _, mapping, template, _, ws = template_case
    notes = _protected_notes(ws)
    assert notes
    assert any('NO SON AUDITABLES' in str(text).upper() for text in notes.values())
    assert not (set(notes) & _refs(mapping, template))


@pytest.mark.parametrize('active', [False, True], ids=['inactive-slots', 'full-slots'])
@pytest.mark.parametrize('param_only', [True, False], ids=['parameters-only', 'with-results'])
def test_plan_preserves_template_notes_and_merged_layout(template_case, active, param_only, tmp_path):
    data, mapping, template, path, _ = template_case
    values = _synthetic_values(mapping, template, active)
    # Jupiter has no resultants even when a caller requests them.
    include_results = not param_only and template.template_group != JUPITER_GROUP
    sources = {'PARAM', 'MOLD', 'MANUAL'} | ({'RESUL'} if include_results else set())
    manual = {'MachineNumber': 'TEST', 'MoldNumber': 'I-9000', 'InjectionControlMode': int(active)}
    plan = PlanBuilder(mapping, data).build(template.template_group, template.layout_id,
                                          values, manual, enabled_sources=sources)
    operations = plan.operations
    if not include_results:
        operations += _param_only_formula_cleanup(path, template_case[-1].title, mapping,
                                                  template.template_group, template.layout_id)
    book = load_workbook(path)
    try:
        ws = book.active
        lookup = _anchors(ws)
        notes = _protected_notes(ws)
        merges = set(map(str, ws.merged_cells.ranges))
        area = str(ws.print_area)
        fills = {op.ref_cell: copy(ws[op.ref_cell].fill) for op in operations if op.kind == 'COPY_FILL'}
        for op in operations:
            assert lookup.get(op.cell, op.cell) == op.cell, (op.kind, op.cell)
            assert op.cell not in notes, (op.kind, op.cell)
            if op.kind == 'WRITE':
                ws[op.cell] = op.value
            elif op.kind == 'CLEAR':
                ws[op.cell] = None
            elif op.kind == 'COPY_FILL':
                ws[op.cell].fill = copy(fills[op.ref_cell])
            else:
                raise AssertionError(op.kind)
        assert _protected_notes(ws) == notes
        assert set(map(str, ws.merged_cells.ranges)) == merges
        assert str(ws.print_area) == area
        if template.template_group == JUPITER_GROUP:
            for col in ('B', 'Z', 'BA', 'BY'):
                assert [ws[f'{col}{r}'].value for r in range(26, 30)] == [1, 2, 3, 'End']
        book.save(tmp_path / 'layout_validation_only.xlsx')
    finally:
        book.close()


def test_geniii_500_old_merge_count_still_fails_strict_validation(catalogs):
    _, mapping = catalogs
    template = next(t for t in mapping.templates if t.layout_id == 'ZERES_GENIII_V1')
    assert template.reference_merged_ranges == 545
    old_definition = replace(template, reference_merged_ranges=544)
    with pytest.raises(TemplateError, match='expected 544, found 545'):
        validate_template(ROOT / 'plantillas' / template.template_file,
                          old_definition, mapping, strict=True)


def test_genv_observations_use_the_revised_merges(catalogs):
    _, mapping = catalogs
    template = next(t for t in mapping.templates if t.layout_id == 'NAVE1_V1')
    book = load_workbook(ROOT / 'plantillas' / template.template_file)
    try:
        ws = book.active
        merges = set(map(str, ws.merged_cells.ranges))
        assert {'U78:AV78', 'U79:AV82'} <= merges
        assert not ({'U79:AV79', 'U80:AV82'} & merges)
        assert ws['U78'].value == 'OBSERVACIONES'
        assert 'NO SON AUDITABLES' in ws['U79'].value
    finally:
        book.close()
