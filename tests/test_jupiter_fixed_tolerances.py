"""Jupiter barrel tolerances are static template data, never XML output."""
from decimal import Decimal
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import openpyxl
import pytest
from openpyxl.utils.cell import range_boundaries

from app.catalogs import DataCatalog, MappingCatalog
from app.jupiter_xml_parser import JupiterXmlParser
from app.resolver import DataResolver

ROOT = Path(__file__).resolve().parents[1]
GROUP = "HAITIAN_JUPITER_TWOSHOT"
LAYOUT = "JUPITER_TWOSHOT_1080_V1"
TOL_RE = re.compile(r"(?:Jupiter\.)?Barrel[12]Zone[1-5](?:UpperTol|LowerTol)")
TOLERANCE_ANCHORS = {
    f"{column}{row}" for column in ("BC", "BH", "BM", "BR", "BW")
    for row in (52, 54, 57, 59)
}


def _write_barrel_xml(path, tolerance=None):
    root = ET.Element("HMI_Data", {"Version": "1.2"})
    group = ET.SubElement(root, "VarGroup", {"Name": "VG_MoldData"})
    for unit in (1, 2):
        for zone in range(1, 6):
            prefix = f"HeatingNozzle{unit}.sv_ZoneRetain{zone}"
            fields = {"rSetValVis": 200 + 10 * unit + zone, "ModeVis": 3 if zone % 2 else 0}
            if tolerance is not None:
                fields.update({"rUpperTolVis": tolerance, "rLowerTolVis": tolerance})
            for field, value in fields.items():
                node = ET.SubElement(group, "Variable", {"Name": f"{prefix}.{field}"})
                ET.SubElement(node, "Value").text = str(value)
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)
    return path


@pytest.mark.parametrize("tolerance", [10, 99, "not-a-number", None])
def test_xml_barrel_tolerances_never_affect_normalized_output(tmp_path, tolerance):
    expected_parser = JupiterXmlParser()
    expected = expected_parser.parse(_write_barrel_xml(tmp_path / "baseline.xml"))
    parser = JupiterXmlParser()
    actual = parser.parse(_write_barrel_xml(tmp_path / "sample.xml", tolerance))
    assert not any(TOL_RE.fullmatch(key) for key in actual)
    assert actual == expected
    assert parser.process_warnings == expected_parser.process_warnings
    # The first five physical barrel zones are always used; ModeVis is not the zone-count authority.
    for unit in (1, 2):
        for zone in range(1, 6):
            prefix = f"Jupiter.Barrel{unit}Zone{zone}"
            assert actual[f"{prefix}Temp"][0].raw_value == Decimal(200 + 10 * unit + zone)
            assert actual[f"{prefix}Mode"][0].raw_value == Decimal(1)


def test_barrel_tolerance_keys_are_removed_from_both_catalogs():
    for filename in ("Data.xlsx", "Mapeo.xlsx"):
        book = openpyxl.load_workbook(ROOT / "data" / filename, read_only=True)
        try:
            for ws in book:
                for row in ws.iter_rows(values_only=True):
                    for value in row:
                        assert not TOL_RE.search(str(value or "")), (filename, ws.title, value)
        finally:
            book.close()
    data = DataCatalog(ROOT / "data/Data.xlsx")
    mapping = MappingCatalog(ROOT / "data/Mapeo.xlsx")
    assert len([r for r in data.param_rules if (r.source_name or "").startswith("Jupiter.")]) == 875
    assert len(mapping.rows_for(GROUP, LAYOUT)) == 882


def _cells(ref):
    c1, r1, c2, r2 = range_boundaries(ref)
    return {(row, col) for row in range(r1, r2 + 1) for col in range(c1, c2 + 1)}


def test_no_jupiter_destination_touches_fixed_tolerance_cells():
    book = openpyxl.load_workbook(ROOT / "plantillas/Haitian Jupiter TwoShot_1080.xlsx")
    mapping = openpyxl.load_workbook(ROOT / "data/Mapeo.xlsx", read_only=True)
    try:
        ws = book.active
        protected = set()
        for anchor in TOLERANCE_ANCHORS:
            assert ws[anchor].value is not None  # retain the existing template constant
            area = next((str(area) for area in ws.merged_cells.ranges if anchor in area), anchor)
            protected.update(_cells(area))
        for sheet_name in ("Mapping", "CoreLayout", "ValveGateLayout"):
            rows = mapping[sheet_name].iter_rows(values_only=True)
            header = next(rows)
            for row in rows:
                item = dict(zip(header, row))
                if item.get("TEMPLATE GROUP") != GROUP or item.get("LAYOUT ID") != LAYOUT:
                    continue
                target = str(item.get("CELL") or "")
                if re.fullmatch(r"\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?", target):
                    assert not (_cells(target) & protected), (sheet_name, item)
    finally:
        mapping.close()
        book.close()


def test_barrel_setpoints_and_states_still_resolve(tmp_path):
    records = JupiterXmlParser().parse(_write_barrel_xml(tmp_path / "barrel.xml", 99))
    data = DataCatalog(ROOT / "data/Data.xlsx")
    required = {f"Barrel{unit}Zone{zone}{suffix}"
                for unit in (1, 2) for zone in range(1, 6) for suffix in ("Temp", "Mode")}
    resolver = DataResolver(data)
    values = resolver.resolve_param(records, required_keys=required)
    assert set(values) == required
    assert not resolver.parameter_warnings
    for unit in (1, 2):
        for zone in range(1, 6):
            assert values[f"Barrel{unit}Zone{zone}Temp"] == Decimal(200 + 10 * unit + zone)
            assert values[f"Barrel{unit}Zone{zone}Mode"] == "ON"
