from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from .errors import ConfigurationError, SourceDataError
from .utils import clean_text


_VALID_MOLD_RE = re.compile(r"^I-\d{4}$", re.IGNORECASE)
_REQUIRED_COLUMNS = (
    "Molde",
    "Numero articulo",
    "Descripción",
    "Art Cliente",
    "Tiempo Ciclo",
    "Cliente",
)


@dataclass(frozen=True)
class MoldRecord:
    mold_number: str
    internal_code: str
    description: str
    customer_part_number: str
    qad_cycle_time: int | None
    customer: str


@dataclass(frozen=True)
class MoldData:
    mold_number: str
    description: str
    internal_code: str
    customer_part_number: str
    qad_cycle_time: int | None
    customer: str
    cavities: str

    def as_values(self) -> dict[str, Any]:
        values: dict[str, Any] = {}
        if self.description:
            values["MoldDescription"] = self.description
        if self.internal_code:
            values["MoldInternalCode"] = self.internal_code
        if self.customer_part_number:
            values["MoldCustomerPartNumber"] = self.customer_part_number
        if self.qad_cycle_time is not None:
            values["MoldQADCycleTime"] = self.qad_cycle_time
        if self.customer:
            values["MoldCustomer"] = self.customer
        if self.cavities:
            values["MoldCavities"] = self.cavities
        return values


class MoldCatalog:
    """Catálogo de información maestra por número de molde.

    El archivo Moldes.xlsx puede contener varias filas para un mismo molde
    (cavidades/artículos distintos). La consolidación es deliberadamente
    conservadora: conserva valores únicos en el orden del archivo, usa el mayor
    Ciclo QAD y solo abrevia descripciones cuando la única diferencia es el
    último token, por ejemplo ``SEAL FRNT DOOR RH / LH``.
    """

    SUPPORTED_KEYS = frozenset(
        {
            "MoldDescription",
            "MoldInternalCode",
            "MoldCustomerPartNumber",
            "MoldQADCycleTime",
            "MoldCustomer",
            "MoldCavities",
        }
    )

    def __init__(self, path: Path):
        self.path = Path(path)
        if not self.path.is_file():
            raise FileNotFoundError(f"No se encontró Moldes.xlsx: {self.path}")
        self._records = self._load_records()
        self._cavities = self._load_cavities()

    @property
    def mold_numbers(self) -> frozenset[str]:
        return frozenset(self._records)

    def resolve(self, mold_number: str) -> MoldData | None:
        key = self.normalize_mold_number(mold_number)
        if not key or key not in self._records:
            return None

        rows = self._records[key]
        descriptions = _unique_nonblank(row.description for row in rows)
        internal_codes = _unique_nonblank(row.internal_code for row in rows)
        customer_parts = _unique_nonblank(row.customer_part_number for row in rows)
        customers = _unique_nonblank(row.customer for row in rows)
        cycle_values = [row.qad_cycle_time for row in rows if row.qad_cycle_time is not None]

        return MoldData(
            mold_number=key,
            description=_combine_descriptions(descriptions),
            internal_code=" / ".join(internal_codes),
            customer_part_number=" / ".join(customer_parts),
            qad_cycle_time=max(cycle_values) if cycle_values else None,
            customer=" / ".join(customers),
            cavities=self._cavities.get(key, ""),
        )

    @staticmethod
    def normalize_mold_number(value: Any) -> str:
        text = clean_text(value).upper()
        return text if _VALID_MOLD_RE.fullmatch(text) else ""

    def _load_records(self) -> dict[str, tuple[MoldRecord, ...]]:
        wb = load_workbook(self.path, data_only=True, read_only=True)
        try:
            if not wb.sheetnames:
                raise ConfigurationError("Moldes.xlsx no contiene hojas")
            ws = wb["Moldes"] if "Moldes" in wb.sheetnames else wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            try:
                headers = tuple(clean_text(value) for value in next(rows))
            except StopIteration as exc:
                raise ConfigurationError("Moldes.xlsx está vacío") from exc

            missing = [name for name in _REQUIRED_COLUMNS if name not in headers]
            if missing:
                raise ConfigurationError(
                    "Moldes.xlsx no contiene las columnas requeridas: " + ", ".join(missing)
                )
            index = {name: headers.index(name) for name in _REQUIRED_COLUMNS}

            grouped: dict[str, list[MoldRecord]] = {}
            seen_rows: set[tuple[Any, ...]] = set()
            for excel_row, values in enumerate(rows, start=2):
                mold = self.normalize_mold_number(values[index["Molde"]])
                if not mold:
                    continue

                internal_code = clean_text(values[index["Numero articulo"]])
                description = _clean_inline_text(values[index["Descripción"]])
                customer_part = clean_text(values[index["Art Cliente"]])
                customer = _clean_inline_text(values[index["Cliente"]])
                cycle = _parse_cycle(values[index["Tiempo Ciclo"]], excel_row=excel_row)

                signature = (
                    mold,
                    internal_code,
                    description,
                    customer_part,
                    cycle,
                    customer,
                )
                if signature in seen_rows:
                    continue
                seen_rows.add(signature)

                grouped.setdefault(mold, []).append(
                    MoldRecord(
                        mold_number=mold,
                        internal_code=internal_code,
                        description=description,
                        customer_part_number=customer_part,
                        qad_cycle_time=cycle,
                        customer=customer,
                    )
                )

            return {mold: tuple(records) for mold, records in grouped.items()}
        finally:
            wb.close()

    def _load_cavities(self) -> dict[str, str]:
        """Carga la hoja CAVS sin convertirla en requisito bloqueante.

        Un molde puede existir solo en la tabla principal o solo en CAVS. En
        ambos casos no se inventa información: únicamente se devuelve una
        cavidad cuando el mismo I-#### existe en ambas fuentes. Los duplicados
        equivalentes de CAVS se consolidan; si existieran valores realmente
        conflictivos para un mismo molde, se deja el campo vacío.
        """

        wb = load_workbook(self.path, data_only=True, read_only=True)
        try:
            if "CAVS" not in wb.sheetnames:
                return {}

            ws = wb["CAVS"]
            rows = ws.iter_rows(values_only=True)
            try:
                headers = tuple(clean_text(value) for value in next(rows))
            except StopIteration:
                return {}

            required = ("Molde", "CAVS")
            if any(name not in headers for name in required):
                return {}
            index = {name: headers.index(name) for name in required}

            grouped: dict[str, list[str]] = {}
            for values in rows:
                mold = self.normalize_mold_number(values[index["Molde"]])
                if not mold or mold not in self._records:
                    continue

                cavities = _normalize_cavities(values[index["CAVS"]])
                if not cavities:
                    continue

                current = grouped.setdefault(mold, [])
                if cavities.casefold() not in {value.casefold() for value in current}:
                    current.append(cavities)

            return {
                mold: values[0]
                for mold, values in grouped.items()
                if len(values) == 1
            }
        finally:
            wb.close()


def _clean_inline_text(value: Any) -> str:
    return " ".join(clean_text(value).split())


def _normalize_cavities(value: Any) -> str:
    text = _clean_inline_text(value)
    if not text:
        return ""
    return re.sub(r"\s*\+\s*", "+", text)


def _parse_cycle(value: Any, *, excel_row: int) -> int | None:
    text = clean_text(value)
    if not text:
        return None
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise SourceDataError(
            f"Moldes.xlsx fila {excel_row}: Tiempo Ciclo no es numérico: {value!r}"
        ) from exc
    if number != number.to_integral_value():
        raise SourceDataError(
            f"Moldes.xlsx fila {excel_row}: Tiempo Ciclo debe ser entero: {value!r}"
        )
    return int(number)


def _unique_nonblank(values: Any) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _clean_inline_text(value)
        if not text:
            continue
        key = text.casefold()
        if key in seen:
            continue
        seen.add(key)
        result.append(text)
    return result


def _combine_descriptions(descriptions: list[str]) -> str:
    if not descriptions:
        return ""
    if len(descriptions) == 1:
        return descriptions[0]

    tokens = [description.split() for description in descriptions]
    first_prefix = tokens[0][:-1]
    if first_prefix and all(parts[:-1] == first_prefix for parts in tokens[1:]):
        suffixes = _unique_nonblank(parts[-1] for parts in tokens)
        if len(suffixes) == len(descriptions):
            return f"{' '.join(first_prefix)} {suffixes[0]}" + " / " + " / ".join(
                suffixes[1:]
            )

    return " / ".join(descriptions)
