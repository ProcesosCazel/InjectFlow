from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .errors import ConfigurationError
from .utils import Table, clean_text, is_truthy, require_columns, split_pipe


def _read_table(ws) -> Table:
    iterator = ws.iter_rows(values_only=True)
    try:
        first = next(iterator)
    except StopIteration:
        return Table((), ())
    headers = tuple(clean_text(v) for v in first)
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if not any(v is not None and clean_text(v) != "" for v in values):
            continue
        row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        rows.append(row)
    return Table(headers, tuple(rows))


@dataclass(frozen=True)
class ParamRule:
    source_name: str | None
    rule_id: str
    key: str
    source: str
    calc_type: str
    factor: Decimal | None
    offset: Decimal
    state_map: str | None
    dependency: tuple[str, ...]
    decimals: int | None
    priority: int | None
    validity_rule: str | None
    active: bool


@dataclass(frozen=True)
class ResulRule:
    source_name: str | None
    rule_id: str
    key: str
    source: str
    calc_type: str
    window: int | None
    round_mode: str
    decimals: int | None
    offset: Decimal
    dependency: tuple[str, ...]
    active: bool


class DataCatalog:
    def __init__(self, path: Path):
        self.path = Path(path)
        try:
            wb = load_workbook(self.path, data_only=False, read_only=True)
        except Exception as exc:
            raise ConfigurationError(f"Cannot open Data.xlsx: {self.path}: {exc}") from exc

        required_sheets = {"Config", "Param", "Resul", "StateMap"}
        missing = required_sheets.difference(wb.sheetnames)
        if missing:
            raise ConfigurationError(f"Data.xlsx is missing sheets: {sorted(missing)}")

        self.config_table = _read_table(wb["Config"])
        self.param_table = _read_table(wb["Param"])
        self.resul_table = _read_table(wb["Resul"])
        self.state_table = _read_table(wb["StateMap"])
        self.machine_constants_table = (
            _read_table(wb["MachineConstants"])
            if "MachineConstants" in wb.sheetnames
            else Table((), ())
        )

        self.config = {
            clean_text(row.get("SETTING")): row.get("VALUE")
            for row in self.config_table.rows
            if clean_text(row.get("SETTING"))
        }
        self.machine_constants: dict[str, dict[str, Any]] = {}
        if self.machine_constants_table.headers:
            require_columns(
                self.machine_constants_table,
                ["MACHINE", "KEY", "VALUE"],
                "MachineConstants",
            )
            for row in self.machine_constants_table.rows:
                machine = clean_text(row.get("MACHINE")).upper()
                key = clean_text(row.get("KEY")).upper()
                if not machine or not key:
                    continue
                bucket = self.machine_constants.setdefault(machine, {})
                value = row.get("VALUE")
                if key in bucket and bucket[key] != value:
                    raise ConfigurationError(
                        f"MachineConstants contains conflicting duplicate {machine}/{key}"
                    )
                bucket[key] = value

        self.state_maps: dict[str, dict[Any, Any]] = {}
        for row in self.state_table.rows:
            map_id = clean_text(row.get("MAP ID"))
            if not map_id:
                continue
            raw = row.get("RAW VALUE")
            if isinstance(raw, float) and raw.is_integer():
                raw = int(raw)
            self.state_maps.setdefault(map_id, {})[raw] = row.get("OUTPUT VALUE")

        self.param_rules = self._parse_param_rules()
        self.resul_rules = self._parse_resul_rules()

    def _parse_param_rules(self) -> tuple[ParamRule, ...]:
        required = ["ID", "PARAMETER KEY", "SOURCE", "CALC TYPE", "ACTIVE"]
        require_columns(self.param_table, required, "Param")
        out: list[ParamRule] = []
        for row in self.param_table.rows:
            active = is_truthy(row.get("ACTIVE"))
            key = clean_text(row.get("PARAMETER KEY"))
            if not key:
                continue
            factor = row.get("FACTOR")
            offset = row.get("OFFSET")
            priority = row.get("PRIORITY")
            decimals = row.get("DECIMALS")
            out.append(
                ParamRule(
                    source_name=clean_text(row.get("MACHINE PARAMETER NAME")) or None,
                    rule_id=clean_text(row.get("ID")),
                    key=key,
                    source=clean_text(row.get("SOURCE")),
                    calc_type=clean_text(row.get("CALC TYPE")).upper(),
                    factor=Decimal(str(factor)) if factor is not None and clean_text(factor) != "" else None,
                    offset=Decimal(str(offset if offset is not None else 0)),
                    state_map=clean_text(row.get("STATE MAP")) or None,
                    dependency=tuple(split_pipe(row.get("DEPENDENCY"))),
                    decimals=int(decimals) if decimals is not None and clean_text(decimals) != "" else None,
                    priority=int(priority) if priority is not None and clean_text(priority) != "" else None,
                    validity_rule=clean_text(row.get("VALIDITY RULE")) or None,
                    active=active,
                )
            )
        return tuple(out)

    def _parse_resul_rules(self) -> tuple[ResulRule, ...]:
        required = ["ID", "RESULTANT KEY", "SOURCE", "CALC TYPE", "ACTIVE"]
        require_columns(self.resul_table, required, "Resul")
        out: list[ResulRule] = []
        for row in self.resul_table.rows:
            key = clean_text(row.get("RESULTANT KEY"))
            if not key:
                continue
            window = row.get("WINDOW")
            decimals = row.get("DECIMALS")
            offset = row.get("OFFSET")
            out.append(
                ResulRule(
                    source_name=clean_text(row.get("MACHINE RESULTANT NAME")) or None,
                    rule_id=clean_text(row.get("ID")),
                    key=key,
                    source=clean_text(row.get("SOURCE")),
                    calc_type=clean_text(row.get("CALC TYPE")).upper(),
                    window=int(window) if window is not None and clean_text(window) != "" else None,
                    round_mode=clean_text(row.get("ROUND MODE")) or "NEAREST",
                    decimals=int(decimals) if decimals is not None and clean_text(decimals) != "" else None,
                    offset=Decimal(str(offset if offset is not None else 0)),
                    dependency=tuple(split_pipe(row.get("DEPENDENCY"))),
                    active=is_truthy(row.get("ACTIVE")),
                )
            )
        return tuple(out)

    def cfg(self, name: str, default: Any = None) -> Any:
        return self.config.get(name, default)


@dataclass(frozen=True)
class MappingRow:
    map_id: str
    template_group: str
    layout_id: str
    key: str
    source: str
    role: str
    cell: str | None
    zone_index: int | None
    controlled_by: str | None
    activation_rule: str
    format_rule: str
    map_status: str
    template_default: str | None
    notes: str | None


@dataclass(frozen=True)
class ZoneRule:
    template_group: str
    layout_id: str
    control_key: str
    control_value: Any
    active_indexes: frozenset[int]
    rule: str
    active_fill_ref: str
    inactive_fill_ref: str


@dataclass(frozen=True)
class TemplateDefinition:
    template_group: str
    template_file: str
    layout_id: str
    sheet_policy: str
    mold_cell: str | None
    print_area: str | None
    write_engine: str
    active: bool
    reference_merged_ranges: int | None
    reference_layout_signature: str | None


class MappingCatalog:
    REQUIRED_SHEETS = {"Mapping", "ZoneRules", "CoreLogic", "CoreLayout", "CoreFields", "Templates", "Machines"}
    VALVE_GATE_SHEETS = {"ValveGateLogic", "ValveGateLayout", "ValveGateFields"}

    def __init__(self, path: Path):
        self.path = Path(path)
        try:
            wb = load_workbook(self.path, data_only=False, read_only=True)
        except Exception as exc:
            raise ConfigurationError(f"Cannot open Mapeo.xlsx: {self.path}: {exc}") from exc
        missing = self.REQUIRED_SHEETS.difference(wb.sheetnames)
        if missing:
            raise ConfigurationError(f"Mapeo.xlsx is missing sheets: {sorted(missing)}")
        self.tables = {name: _read_table(wb[name]) for name in self.REQUIRED_SHEETS}
        for name in self.VALVE_GATE_SHEETS:
            if name in wb.sheetnames:
                self.tables[name] = _read_table(wb[name])
        if "ManualInputs" in wb.sheetnames:
            self.tables["ManualInputs"] = _read_table(wb["ManualInputs"])

        self.mapping_rows = self._parse_mapping()
        has_valve_dynamic = any(r.map_status == "VALVE_GATE_DYNAMIC" for r in self.mapping_rows)
        missing_valve = self.VALVE_GATE_SHEETS.difference(self.tables)
        if has_valve_dynamic and missing_valve:
            raise ConfigurationError(
                "Mapeo.xlsx contains VALVE_GATE_DYNAMIC rows but is missing sheets: "
                f"{sorted(missing_valve)}"
            )
        self.zone_rules = self._parse_zone_rules()
        self.templates = self._parse_templates()
        self.machines, self.machine_layouts = self._parse_machines()
        self.core_logic_rows = self.tables["CoreLogic"].rows
        self.core_layout_rows = self.tables["CoreLayout"].rows
        self.core_field_rows = self.tables["CoreFields"].rows
        self.valve_gate_logic_rows = self.tables.get("ValveGateLogic", Table((), ())).rows
        self.valve_gate_layout_rows = self.tables.get("ValveGateLayout", Table((), ())).rows
        self.valve_gate_field_rows = self.tables.get("ValveGateFields", Table((), ())).rows

    def _parse_mapping(self) -> tuple[MappingRow, ...]:
        table = self.tables["Mapping"]
        required = ["MAP ID", "TEMPLATE GROUP", "LAYOUT ID", "KEY", "SOURCE", "ROLE", "CELL", "FORMAT RULE", "MAP STATUS"]
        require_columns(table, required, "Mapping")
        out: list[MappingRow] = []
        for row in table.rows:
            idx = row.get("ZONE INDEX")
            try:
                zone_index = int(idx) if idx is not None and clean_text(idx) != "" and str(idx).strip().lstrip("-").isdigit() else None
            except (TypeError, ValueError):
                zone_index = None
            out.append(
                MappingRow(
                    map_id=clean_text(row.get("MAP ID")),
                    template_group=clean_text(row.get("TEMPLATE GROUP")),
                    layout_id=clean_text(row.get("LAYOUT ID")),
                    key=clean_text(row.get("KEY")),
                    source=clean_text(row.get("SOURCE")),
                    role=clean_text(row.get("ROLE")).upper(),
                    cell=(clean_text(row.get("CELL")) or None),
                    zone_index=zone_index,
                    controlled_by=clean_text(row.get("CONTROLLED BY")) or None,
                    activation_rule=clean_text(row.get("ACTIVATION RULE")),
                    format_rule=clean_text(row.get("FORMAT RULE")).upper(),
                    map_status=clean_text(row.get("MAP STATUS")).upper(),
                    template_default=clean_text(row.get("TEMPLATE DEFAULT")) or None,
                    notes=clean_text(row.get("NOTES")) or None,
                )
            )
        return tuple(out)

    def _parse_zone_rules(self) -> tuple[ZoneRule, ...]:
        table = self.tables["ZoneRules"]
        out: list[ZoneRule] = []
        for row in table.rows:
            control = clean_text(row.get("CONTROL KEY"))
            active_text = clean_text(row.get("ACTIVE INDEXES"))
            if not control or "{n}" in control or active_text.lower() == "same zone":
                continue
            indexes = frozenset(int(x.strip()) for x in active_text.split(",") if x.strip())
            value = row.get("CONTROL VALUE")
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            out.append(
                ZoneRule(
                    template_group=clean_text(row.get("TEMPLATE GROUP")),
                    layout_id=clean_text(row.get("LAYOUT ID")),
                    control_key=control,
                    control_value=value,
                    active_indexes=indexes,
                    rule=clean_text(row.get("RULE")),
                    active_fill_ref=clean_text(row.get("ACTIVE FILL REF")),
                    inactive_fill_ref=clean_text(row.get("INACTIVE FILL REF")),
                )
            )
        return tuple(out)

    def _parse_templates(self) -> tuple[TemplateDefinition, ...]:
        out: list[TemplateDefinition] = []
        for row in self.tables["Templates"].rows:
            merged = row.get("REFERENCE MERGED RANGES")
            out.append(
                TemplateDefinition(
                    template_group=clean_text(row.get("TEMPLATE GROUP")),
                    template_file=clean_text(row.get("TEMPLATE FILE")),
                    layout_id=clean_text(row.get("LAYOUT ID")),
                    sheet_policy=clean_text(row.get("SHEET POLICY")),
                    mold_cell=clean_text(row.get("MOLD CELL")) or None,
                    print_area=clean_text(row.get("PRINT AREA")) or None,
                    write_engine=clean_text(row.get("WRITE ENGINE")) or "EXCEL_COM",
                    active=is_truthy(row.get("ACTIVE")),
                    reference_merged_ranges=int(merged) if merged is not None and clean_text(merged) != "" else None,
                    reference_layout_signature=clean_text(row.get("REFERENCE LAYOUT SIGNATURE")) or None,
                )
            )
        return tuple(out)

    def _parse_machines(self) -> tuple[dict[str, str], dict[str, str]]:
        """Parse active machines and their optional machine-specific layout.

        Older Mapeo.xlsx files identify a machine only by TEMPLATE GROUP. The current
        implementation keeps that behavior for backward compatibility, but when Machines has a
        LAYOUT ID it becomes authoritative for selecting the template inside a
        shared family such as HAITIAN_ZE_III.
        """
        groups: dict[str, str] = {}
        layouts: dict[str, str] = {}
        for row in self.tables["Machines"].rows:
            if not is_truthy(row.get("ACTIVE")):
                continue
            machine = clean_text(row.get("MACHINE")).upper()
            group = clean_text(row.get("TEMPLATE GROUP"))
            layout_id = clean_text(row.get("LAYOUT ID"))
            if not machine:
                continue
            if not group:
                raise ConfigurationError(
                    f"Machine {machine!r} is active but has no TEMPLATE GROUP in Mapeo.xlsx"
                )
            groups[machine] = group
            if layout_id:
                layouts[machine] = layout_id
        return groups, layouts

    def resolve_machine(
        self, machine: str, *, allow_inactive: bool = False
    ) -> tuple[str, TemplateDefinition]:
        key = machine.strip().upper()
        if allow_inactive:
            # Used only by an explicit plan-only development run. Never changes
            # the catalog, its active-machine dictionary, or the desktop list.
            rows = [r for r in self.tables["Machines"].rows
                    if clean_text(r.get("MACHINE")).upper() == key]
            if len(rows) != 1:
                raise ConfigurationError(f"Expected one machine definition for {key}, found {len(rows)}")
            group = clean_text(rows[0].get("TEMPLATE GROUP"))
            layout_id = clean_text(rows[0].get("LAYOUT ID"))
            candidates = [t for t in self.templates if t.template_group == group
                          and (not layout_id or t.layout_id == layout_id)]
            if not group or len(candidates) != 1:
                raise ConfigurationError(f"Expected one template definition for {key}/{group}/{layout_id}")
            return group, candidates[0]
        if key not in self.machines:
            raise ConfigurationError(
                f"Machine {machine!r} is not active in Mapeo.xlsx. Valid machines: {', '.join(sorted(self.machines))}"
            )

        group = self.machines[key]
        layout_id = self.machine_layouts.get(key)

        if layout_id:
            candidates = [
                t for t in self.templates
                if t.template_group == group
                and t.layout_id == layout_id
                and t.active
            ]
            if len(candidates) != 1:
                raise ConfigurationError(
                    f"Expected one active template for machine {key} using "
                    f"{group}/{layout_id}, found {len(candidates)}"
                )
            return group, candidates[0]

        # Backward-compatible fallback for legacy mapping files that do not
        # define LAYOUT ID per machine.  This remains valid only while the
        # template group has exactly one active template.
        candidates = [t for t in self.templates if t.template_group == group and t.active]
        if len(candidates) != 1:
            raise ConfigurationError(
                f"Machine {key} has no LAYOUT ID and template group {group} has "
                f"{len(candidates)} active templates. Define LAYOUT ID in Machines."
            )
        return group, candidates[0]

    def rows_for(self, group: str, layout_id: str) -> tuple[MappingRow, ...]:
        return tuple(r for r in self.mapping_rows if r.template_group == group and r.layout_id == layout_id)

    def zone_rule(self, group: str, layout_id: str, control_key: str, control_value: Any) -> ZoneRule:
        def norm(v: Any) -> str:
            if isinstance(v, Decimal):
                if v == v.to_integral_value():
                    return str(int(v))
                return str(v.normalize())
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return clean_text(v).upper()
        target = norm(control_value)
        matches = [
            r for r in self.zone_rules
            if r.template_group == group and r.layout_id == layout_id and r.control_key == control_key and norm(r.control_value) == target
        ]
        if len(matches) != 1:
            raise ConfigurationError(
                f"No unique ZoneRules entry for {control_key}={control_value!r} in {group}/{layout_id}"
            )
        return matches[0]

    def core_settings(self, group: str, layout_id: str) -> dict[str, Any]:
        rows = [r for r in self.core_logic_rows if clean_text(r.get("TEMPLATE GROUP")) == group and clean_text(r.get("LAYOUT ID")) == layout_id]
        return {clean_text(r.get("SETTING")): r.get("VALUE") for r in rows if clean_text(r.get("SETTING"))}

    def core_layout(self, group: str, layout_id: str) -> tuple[dict[str, Any], ...]:
        return tuple(r for r in self.core_layout_rows if clean_text(r.get("TEMPLATE GROUP")) == group and clean_text(r.get("LAYOUT ID")) == layout_id)

    def core_fields(self, group: str, layout_id: str) -> tuple[dict[str, Any], ...]:
        return tuple(r for r in self.core_field_rows if clean_text(r.get("TEMPLATE GROUP")) == group and clean_text(r.get("LAYOUT ID")) == layout_id)

    def valve_gate_settings(self, group: str, layout_id: str) -> dict[str, Any]:
        rows = [
            r for r in self.valve_gate_logic_rows
            if clean_text(r.get("TEMPLATE GROUP")) == group
            and clean_text(r.get("LAYOUT ID")) == layout_id
        ]
        return {
            clean_text(r.get("SETTING")): r.get("VALUE")
            for r in rows
            if clean_text(r.get("SETTING"))
        }

    def valve_gate_layout(self, group: str, layout_id: str) -> tuple[dict[str, Any], ...]:
        return tuple(
            r for r in self.valve_gate_layout_rows
            if clean_text(r.get("TEMPLATE GROUP")) == group
            and clean_text(r.get("LAYOUT ID")) == layout_id
        )

    def valve_gate_fields(self, group: str, layout_id: str) -> tuple[dict[str, Any], ...]:
        return tuple(
            r for r in self.valve_gate_field_rows
            if clean_text(r.get("TEMPLATE GROUP")) == group
            and clean_text(r.get("LAYOUT ID")) == layout_id
        )


    def manual_input_keys(self, group: str | None = None) -> set[str]:
        """Return manual input keys declared in the optional ManualInputs sheet."""
        table = self.tables.get("ManualInputs")
        if table is None:
            return set()
        result: set[str] = set()
        for row in table.rows:
            row_group = clean_text(row.get("TEMPLATE GROUP"))
            if group is not None and row_group and row_group != group:
                continue
            key = clean_text(row.get("INPUT KEY"))
            if key:
                result.add(key)
        return result

    def required_keys_for(
        self,
        group: str,
        layout_id: str,
        enabled_sources: set[str] | None = None,
    ) -> set[str]:
        """Return logical Data.xlsx keys needed to generate this layout.

        ``enabled_sources`` filters physical mapping rows by their SOURCE column.
        This is used by the ``Solo Param.dat`` mode to exclude resultants without
        weakening the requirements for Param.dat-driven zones, states or cores.
        Core selectors/fields remain required because the current core model is
        sourced from Param.dat.
        """
        normalized_sources = (
            {clean_text(source).upper() for source in enabled_sources}
            if enabled_sources is not None
            else None
        )
        required: set[str] = set()
        for row in self.rows_for(group, layout_id):
            row_source = clean_text(row.source).upper()
            if normalized_sources is not None and row_source not in normalized_sources:
                continue
            if row.map_status in {"MAPPED", "CORE_DYNAMIC", "VALVE_GATE_DYNAMIC"} and row.key:
                required.add(row.key)
            if row.controlled_by:
                # Some v1.3 rules have more than one controller, separated by |
                # (for example InjectionControlMode|InjectionZones).
                required.update(split_pipe(row.controlled_by))
        settings = self.core_settings(group, layout_id)
        required.update(split_pipe(settings.get("SELECTOR_KEYS")))
        for row in self.core_fields(group, layout_id):
            status = clean_text(row.get("STATUS")).upper()
            key = clean_text(row.get("DATA KEY"))
            if key and status in {"MAPPED", "CONTROL"}:
                required.add(key)
        for row in self.valve_gate_fields(group, layout_id):
            status = clean_text(row.get("STATUS")).upper()
            key = clean_text(row.get("DATA KEY"))
            if key and status in {"MAPPED", "CONTROL"}:
                required.add(key)
        return required
