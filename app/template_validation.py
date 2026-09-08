from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .catalogs import MappingCatalog, TemplateDefinition
from .errors import TemplateError


@dataclass(frozen=True)
class TemplateValidationResult:
    sheet_name: str
    merged_ranges: int
    print_area: str | None
    warnings: tuple[str, ...]


def validate_template(
    path: Path,
    template: TemplateDefinition,
    mapping: MappingCatalog,
    *,
    strict: bool = True,
) -> TemplateValidationResult:
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        raise TemplateError(f"Cannot inspect template {path}: {exc}") from exc

    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    if len(visible) != 1:
        raise TemplateError(f"Expected exactly one visible worksheet; found {len(visible)}")
    ws = visible[0]
    warnings: list[str] = []

    merged_count = len(ws.merged_cells.ranges)
    if template.reference_merged_ranges is not None and merged_count != template.reference_merged_ranges:
        msg = f"Merged range count changed: expected {template.reference_merged_ranges}, found {merged_count}"
        if strict:
            raise TemplateError(msg)
        warnings.append(msg)

    actual_print_area = str(ws.print_area) if ws.print_area else None
    expected = (template.print_area or "").replace("$", "").replace("'", "")
    actual_cmp = (actual_print_area or "").replace("$", "").replace("'", "")
    if expected and expected not in actual_cmp:
        warnings.append(f"Print area differs: expected {template.print_area}, found {actual_print_area}")

    # Ensure every physically mapped cell is a normal cell or the top-left anchor of a merged range.
    merged_lookup: dict[str, str] = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col).coordinate
        for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
            for cell in row:
                merged_lookup[cell.coordinate] = anchor
    bad: list[str] = []
    for row in mapping.rows_for(template.template_group, template.layout_id):
        if row.map_status != "MAPPED" or not row.cell or row.cell.upper() in {"NOT USED", "CORE_DYNAMIC", "INPUT_ONLY"}:
            continue
        anchor = merged_lookup.get(row.cell)
        if anchor and anchor != row.cell:
            bad.append(f"{row.map_id}:{row.cell}->{anchor}")
    if bad:
        raise TemplateError(f"Mapped cells are not merged-range anchors: {bad[:10]}")

    return TemplateValidationResult(ws.title, merged_count, actual_print_area, tuple(warnings))
